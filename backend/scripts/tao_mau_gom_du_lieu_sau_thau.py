#!/usr/bin/env python3
"""Sinh biểu mẫu gom DỮ LIỆU SAU ĐẤU THẦU — miếng 0 (chốt 25/08/2026).

Vì sao có file này: workbook `database/database web.xlsx` (03/08/2026) có ba
sheet sau thầu — HOP_DONG · KET_QUA_THAU · GOI_THAU_TIMELINE — đều TRỐNG và đều
trỏ vào mô hình TRƯỚC v3. Điền vào đó thì dữ liệu rơi vào hư không:
`GOI_THAU_TIMELINE` đổ vào `goi_thau_moc` (một trong ba bảng đã chết),
`KET_QUA_THAU` là hình dạng của `goi_thau_ket_qua_ma` (cũng đã chết, và trong v3
kết quả thầu do web TỰ SINH từ số PĐD gõ trên bảng Tổng hợp), còn `HOP_DONG` có
cột `tran_hop_dong` — trái QĐ 17/08 "hợp đồng không có cột giá".

Biểu mẫu này chỉ gom ĐÚNG HAI THỨ mà web không tự sinh ra được:
hợp đồng, và từng lần hàng về.

Quyết định của chủ dự án 21/08 và 25/08/2026 đã áp vào thiết kế:

  • KHÔNG có cột giá / đơn giá / trần hợp đồng          (QĐ 17/08, xác nhận 21/08)
  • Giao hàng ghi theo TỪNG LẦN GIAO, không phải
    ảnh chụp tồn kho định kỳ                             (21/08)
  • Mỗi nhà thầu MỘT hợp đồng → tách hai sheet:
    HOP_DONG và HOP_DONG_MA_HANG                         (25/08)
  • KHÔNG lưu số lô và hạn dùng                          (25/08)
  • Hàng về KHO trước rồi kho mới cấp cho khoa, nên cột
    `khoa` ở GIAO_HANG ĐỂ TRỐNG ĐƯỢC. Phần kho→khoa lấy
    từ lịch sử xuất kho HIS đã nạp, không gõ lại         (25/08)

Dùng lại bộ định dạng của `tao_mau_du_lieu_benh_vien.py` để hai biểu mẫu nhìn
giống nhau: cùng màu theo mức độ bắt buộc, cùng ghi chú trong ô tiêu đề, cùng
kiểu kiểm tra dữ liệu, cùng cách tô đỏ dòng thiếu trường bắt buộc.

    cd backend
    .venv/bin/python scripts/tao_mau_gom_du_lieu_sau_thau.py \\
        ../database/MAU_GOM_DU_LIEU_SAU_THAU.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from tao_mau_du_lieu_benh_vien import (
    BLUE, GRAY, GREEN, NAVY, ORANGE, TEXT, WHITE, YELLOW,
    Field, SheetSpec, add_data_sheet, f, fill, style_header, style_title,
)

DATE = "YYYY-MM-DD"
BOOL = "TRUE | FALSE"

# Mã gói con — khớp GOI_CON trong `frontend/src/features/KhungGoiThau.jsx`.
# Gói bổ sung đã mã hoá luôn tháng mốc trong mã (bs-t1 · bs-t5 · bs-t9) nên
# `goi_con` + `nam` là đủ để tìm đúng DOT_GOI, không cần cột tháng riêng.
GOI_CON = ("18t-dung-chung | 18t-gmhs | 18t-rhm | 18t-tim-mach | 18t-ctch-ntk "
           "| bs-t1 | bs-t5 | bs-t9")


SHEETS: tuple[SheetSpec, ...] = (
    SheetSpec(
        "HOP_DONG",
        "Hợp đồng — mỗi nhà thầu một dòng",
        "Một dòng cho một hợp đồng. Một gói con trúng nhiều nhà thầu thì có nhiều dòng.",
        "hop_dong_v3 (bảng dựng ở miếng 3)",
        "Bắt buộc",
        (
            f("goi_con", "Bắt buộc", "Text",
              "Gói con mà hợp đồng này thuộc về. Dùng cùng `nam` để tìm đúng đợt.",
              "hop_dong_v3.dot_goi_id", allowed=GOI_CON, example="18t-dung-chung"),
            f("nam", "Bắt buộc", "Integer",
              "Năm của đợt đề xuất, KHÔNG phải năm ký hợp đồng. Gói 18T 1/2027–6/2028 thì ghi 2027.",
              "hop_dong_v3.dot_goi_id", allowed=">= 2024", example="2027"),
            f("so_hop_dong", "Bắt buộc", "Text",
              "Số hợp đồng theo văn bản. Duy nhất trong một đợt × gói con. Giữ nguyên dạng text, kể cả khi có dấu / hoặc số 0 đầu.",
              "hop_dong_v3.so_hop_dong", example="01/2027/HĐ-VTYT"),
            f("nha_cung_cap", "Bắt buộc", "Text",
              "Tên nhà thầu đúng như trên hợp đồng.",
              "hop_dong_v3.nha_cung_cap", example="Công ty TNHH Thiết bị Y tế ABC"),
            f("ngay_ky", "Bắt buộc", "Date",
              "Ngày ký hợp đồng.",
              "hop_dong_v3.ngay_ky", allowed=DATE, example="2027-01-15"),
            f("ngay_het_han", "Khuyến nghị", "Date",
              "Ngày hết hiệu lực hợp đồng. Không có sẵn thì để trống, đừng đoán.",
              "hop_dong_v3.ngay_het_han", allowed=DATE, example="2028-07-15"),
            f("co_tuy_chon_30", "Tùy chọn", "Boolean",
              "Hợp đồng này có điều khoản mua thêm 30% không. Không rõ thì để trống.",
              "hop_dong_v3.co_tuy_chon_30", allowed=BOOL, example="TRUE"),
            f("ghi_chu", "Tùy chọn", "Text dài",
              "Bất cứ điều gì cần nói thêm. Không có chỗ nào khác thì ghi vào đây.",
              "hop_dong_v3.ghi_chu"),
        ),
    ),
    SheetSpec(
        "HOP_DONG_MA_HANG",
        "Mã hàng trong từng hợp đồng",
        "Một dòng cho một mã hàng thuộc một hợp đồng. Đây là số CAM KẾT MUA theo hợp đồng.",
        "hop_dong_ma_hang (bảng dựng ở miếng 3)",
        "Bắt buộc",
        (
            f("so_hop_dong", "Bắt buộc", "Text",
              "Phải trùng đúng chữ với một dòng ở sheet HOP_DONG. Sai một ký tự là dòng này mồ côi.",
              "hop_dong_ma_hang.hop_dong_id", example="01/2027/HĐ-VTYT"),
            f("ma_hang", "Bắt buộc", "Text",
              "Mã SKU trong HIS. GIỮ SỐ 0 ĐẦU — định dạng ô là Text trước khi dán.",
              "hop_dong_ma_hang.ma_hang", example="0066114"),
            f("so_luong_hop_dong", "Bắt buộc", "Integer",
              "Số lượng cam kết mua theo hợp đồng cho mã này. KHÔNG phải số đã giao.",
              "hop_dong_ma_hang.so_luong_hop_dong", allowed=">= 0", example="980"),
            f("dvt", "Khuyến nghị", "Text",
              "Đơn vị tính ghi trên hợp đồng. Lệch với ĐVT trong danh mục vật tư thì web sẽ cảnh báo chứ không tự quy đổi.",
              "hop_dong_ma_hang.dvt", example="Đôi"),
            f("ghi_chu", "Tùy chọn", "Text dài",
              "Ví dụ: mã này chia làm hai đợt giao, hoặc thay thế cho mã khác.",
              "hop_dong_ma_hang.ghi_chu"),
        ),
    ),
    SheetSpec(
        "GIAO_HANG",
        "Từng lần hàng về",
        "Một dòng cho MỘT LẦN giao. Hàng về làm sáu lần thì sáu dòng, không cộng gộp lại.",
        "giao_hang (bảng dựng ở miếng 3)",
        "Bắt buộc",
        (
            f("ngay_giao", "Bắt buộc", "Date",
              "Ngày hàng thực nhận. Đây cũng là mốc bắt đầu đếm cam kết 20/50/80.",
              "giao_hang.ngay_giao", allowed=DATE, example="2027-03-04"),
            f("so_hop_dong", "Bắt buộc", "Text",
              "Hợp đồng mà lần giao này thuộc về. Phải trùng đúng chữ với sheet HOP_DONG.",
              "giao_hang.hop_dong_id", example="01/2027/HĐ-VTYT"),
            f("ma_hang", "Bắt buộc", "Text",
              "Mã SKU trong HIS. GIỮ SỐ 0 ĐẦU.",
              "giao_hang.ma_hang", example="0066114"),
            f("so_luong_thuc_nhan", "Bắt buộc", "Integer",
              "Số thực nhận lần này. Giao thiếu thì ghi đúng số thiếu, đừng ghi số trên phiếu.",
              "giao_hang.so_luong_thuc_nhan", allowed=">= 0", example="200"),
            f("khoa", "Tùy chọn", "Text",
              "ĐỂ TRỐNG nếu hàng về kho chung (đây là trường hợp thường gặp — QĐ 25/08). "
              "Chỉ điền khi nhà thầu giao thẳng cho một khoa. Phần kho cấp cho khoa lấy "
              "từ lịch sử xuất kho HIS, KHÔNG gõ lại ở đây.",
              "giao_hang.khoa", example="Khoa Ngoại tổng hợp"),
            f("ghi_chu", "Tùy chọn", "Text dài",
              "Ví dụ: giao bù đợt trước, hàng đổi mã, biên bản số mấy.",
              "giao_hang.ghi_chu"),
        ),
    ),
)


def add_guide(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "00_HUONG_DAN"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 118

    def dong(txt: str = "", *, dam: bool = False, mau: str | None = None,
             nghieng: bool = False, cao: int | None = None):
        r = ws.max_row + 1 if ws.max_row > 1 or ws["A1"].value else 1
        c = ws.cell(r, 1, txt)
        c.font = Font(bold=dam, italic=nghieng, color=TEXT)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if mau:
            c.fill = fill(mau)
        if cao:
            ws.row_dimensions[r].height = cao
        return r

    ws["A1"] = "GOM DỮ LIỆU SAU ĐẤU THẦU — hợp đồng và giao hàng"
    style_title(ws["A1"])
    ws.row_dimensions[1].height = 34

    dong()
    dong("Biểu mẫu này chỉ hỏi HAI thứ mà web không tự biết được: hợp đồng ký với ai, "
         "và mỗi lần hàng về bao nhiêu. Mọi con số khác — số trúng, đã chia về khoa, "
         "còn thiếu, tiến độ 20/50/80 — web tự tính, đừng gõ lại.", nghieng=True, cao=44)
    dong()

    dong("  ĐIỀN THEO THỨ TỰ NÀY  ", dam=True, mau=NAVY)
    ws.cell(ws.max_row, 1).font = Font(bold=True, color=WHITE)
    dong()
    for i, (ten, mo_ta) in enumerate((
        ("HOP_DONG", "Mỗi nhà thầu một dòng. Điền hết sheet này trước."),
        ("HOP_DONG_MA_HANG", "Mã hàng thuộc từng hợp đồng. Cột `so_hop_dong` phải trùng đúng chữ với sheet trên."),
        ("GIAO_HANG", "Mỗi lần hàng về một dòng. Điền dần theo thời gian, không cần chờ đủ."),
    ), 1):
        dong(f"    {i}. {ten} — {mo_ta}")
    dong()

    dong("  MÀU Ở DÒNG TIÊU ĐỀ  ", dam=True, mau=NAVY)
    ws.cell(ws.max_row, 1).font = Font(bold=True, color=WHITE)
    dong()
    for mau, nhan in ((ORANGE, "Bắt buộc — thiếu là dòng đó không dùng được"),
                      (YELLOW, "Có điều kiện"),
                      (BLUE, "Khuyến nghị — có thì tốt, không có thì để trống"),
                      (GRAY, "Tùy chọn")):
        r = dong(f"    {nhan}")
        ws.cell(r, 1).fill = fill(mau)
    dong()
    dong("Dòng nào đã gõ dữ liệu mà thiếu trường bắt buộc đầu tiên sẽ TỰ TÔ ĐỎ. "
         "Rê chuột lên ô tiêu đề để đọc giải thích đầy đủ của từng cột; "
         "sheet 01_TU_DIEN_COT liệt kê lại toàn bộ.", cao=32)
    dong()

    dong("  BỐN ĐIỀU ĐÃ CHỐT, ĐỪNG THÊM CỘT  ", dam=True, mau=NAVY)
    ws.cell(ws.max_row, 1).font = Font(bold=True, color=WHITE)
    dong()
    for txt in (
        "KHÔNG có cột giá, đơn giá hay trần hợp đồng. Web này không quản lý tiền "
        "(QĐ 17/08/2026, xác nhận lại 21/08).",
        "Giao hàng ghi TỪNG LẦN GIAO. Không ghi kiểu ảnh chụp tồn kho mỗi tháng — "
        "cách đó nặng gấp mười lần mà 90% số dòng lặp lại y nguyên tháng trước (QĐ 21/08).",
        "KHÔNG ghi số lô và hạn dùng (QĐ 25/08). Cần truy lô thì tra ở phần mềm kho.",
        "Cột `khoa` ở GIAO_HANG ĐỂ TRỐNG khi hàng về kho chung — đây là trường hợp "
        "thường gặp. Phần kho cấp cho từng khoa đã có sẵn trong lịch sử xuất kho HIS "
        "(141.623 dòng đã nạp), không gõ lại (QĐ 25/08).",
    ):
        dong(f"    •  {txt}", cao=32)
    dong()

    dong("  ĐỪNG ĐIỀN BA SHEET NÀY CỦA FILE CŨ  ", dam=True, mau=NAVY)
    ws.cell(ws.max_row, 1).font = Font(bold=True, color=WHITE)
    dong()
    dong("Trong `database/database web.xlsx` còn ba sheet sau thầu đã LẠC HẬU. "
         "Điền vào đó thì dữ liệu không vào được hệ:", cao=30)
    for ten, ly_do in (
        ("KET_QUA_THAU", "trong v3, kết quả thầu do web TỰ SINH từ số PĐD gõ trên bảng Tổng hợp"),
        ("GOI_THAU_TIMELINE", "đổ vào bảng `goi_thau_moc` đã chết từ 17/08"),
        ("HOP_DONG (bản cũ)", "có cột trần hợp đồng, trái quyết định bỏ mọi cột giá; và không neo được vào đợt"),
    ):
        dong(f"    ✗  {ten} — {ly_do}", cao=28)
    dong()

    dong("  GỬI LẠI THẾ NÀO  ", dam=True, mau=NAVY)
    ws.cell(ws.max_row, 1).font = Font(bold=True, color=WHITE)
    dong()
    dong("Điền được tới đâu gửi tới đó, không cần chờ đủ. Trước khi gửi, chạy phép "
         "kiểm để biết dữ liệu có dùng được không:", cao=30)
    r = dong("    cd backend && .venv/bin/python scripts/kiem_mau_gom_du_lieu.py "
             "../database/MAU_GOM_DU_LIEU_SAU_THAU.xlsx --xac-nhan-staging")
    ws.cell(r, 1).font = Font(name="Consolas", color=TEXT)
    ws.cell(r, 1).fill = fill(GREEN)
    dong()
    dong("Phép kiểm nói ngay: mã hàng nào không có trong danh mục, hợp đồng nào bị "
         "mồ côi, gói con nào ghi sai, ngày nào vô lý. Sửa xong chạy lại.", cao=30)


def add_dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("01_TU_DIEN_COT")
    ws.sheet_view.showGridLines = False
    ws.append(("Sheet", "Mục đích", "Cột", "Mức độ", "Kiểu dữ liệu",
               "Giá trị hợp lệ", "Đích database", "Giải thích chi tiết", "Ví dụ"))
    for cell in ws[1]:
        style_header(cell)
    for spec in SHEETS:
        for field in spec.fields:
            ws.append((spec.name, spec.purpose, field.name, field.level, field.dtype,
                       field.allowed, field.target, field.description, field.example))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E2E8F0"))
        row[3].fill = fill({"Bắt buộc": ORANGE, "Có điều kiện": YELLOW,
                            "Khuyến nghị": BLUE, "Tùy chọn": GRAY}.get(row[3].value, WHITE))
    for idx, width in enumerate((22, 46, 22, 15, 14, 40, 40, 72, 26), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"


def build(output: Path) -> None:
    wb = Workbook()
    add_guide(wb)
    add_dictionary(wb)
    for spec in SHEETS:
        add_data_sheet(wb, spec)
    wb.save(output)

    # Mở lại: xác nhận file hợp lệ và header đúng như khai báo.
    check = load_workbook(output, read_only=True)
    thieu = {"00_HUONG_DAN", "01_TU_DIEN_COT", *(s.name for s in SHEETS)} - set(check.sheetnames)
    if thieu:
        raise RuntimeError(f"Workbook thiếu sheet: {sorted(thieu)}")
    for spec in SHEETS:
        thuc = tuple(c.value for c in next(check[spec.name].iter_rows(min_row=1, max_row=1)))
        if thuc != tuple(f_.name for f_ in spec.fields):
            raise RuntimeError(f"Header sheet {spec.name} không khớp: {thuc}")
    check.close()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("output", nargs="?", type=Path,
                    default=Path("MAU_GOM_DU_LIEU_SAU_THAU.xlsx"))
    args = ap.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
