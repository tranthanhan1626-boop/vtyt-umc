#!/usr/bin/env python3
"""Tạo workbook full workflow đã đối chiếu hai file production của bệnh viện."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tao_mau_du_lieu_benh_vien import (
    BLUE,
    GRAY,
    GREEN,
    LIGHT,
    NAVY,
    ORANGE,
    RED,
    TEXT,
    WHITE,
    YELLOW,
    build,
    fill,
    style_header,
    style_title,
)


USAGE_NAME = "so luong su dung full.xlsx"
CATALOG_NAME = "thong tin vat tu y te tieu hao.xlsx"


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_catalog(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [clean(v) for v in next(ws.iter_rows(values_only=True))]
    by_code: dict[str, list[dict]] = defaultdict(list)
    nulls = Counter()
    rows = 0
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        row = dict(zip(headers, values))
        code = clean(row.get("Mã hàng"))
        if not code:
            continue
        rows += 1
        by_code[code].append(row)
        for header in headers:
            if row.get(header) in (None, ""):
                nulls[header] += 1
    wb.close()
    return headers, by_code, rows, nulls


def scan_usage_and_fill(path: Path, target_ws):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [clean(v) for v in next(ws.iter_rows(values_only=True))]
    ix = {name: i for i, name in enumerate(headers)}

    rows = 0
    units = Counter()
    warehouses = set()
    codes: dict[str, dict] = {}
    manager_names: dict[str, str] = {}
    missing_manager_rows = 0
    dates: list[date] = []
    fact_keys = set()
    duplicate_fact_keys = 0
    nulls = Counter()

    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        rows += 1
        don_vi = clean(values[ix["Đơn vị"]])
        kho_xuat = clean(values[ix["Kho xuất"]])
        ma_quan_ly = clean(values[ix["Mã quản lý"]])
        ten_quan_ly = clean(values[ix["Tên quản lý"]])
        ma_hang = clean(values[ix["Mã hàng"]])
        ten_vat_tu = clean(values[ix["Tên vật tư"]])
        dvt = clean(values[ix["ĐVT"]])
        ngay = values[ix["Ngày"]]
        so_luong = values[ix["Số lượng"]]

        if isinstance(ngay, datetime):
            ngay = ngay.date()
        if isinstance(ngay, date):
            nam, thang = ngay.year, ngay.month
            dates.append(ngay)
        else:
            nam = int(values[ix["Ngày - Year"]])
            thang = int(clean(values[ix["Tháng"]]).replace("Tháng", "").strip())

        target_ws.append(
            (
                don_vi,
                kho_xuat,
                ma_hang,
                nam,
                thang,
                so_luong,
                USAGE_NAME,
                None,
                None,
            )
        )

        units[don_vi] += 1
        warehouses.add(kho_xuat)
        if ma_hang not in codes:
            codes[ma_hang] = {
                "ma_hang": ma_hang,
                "ten_vat_tu": ten_vat_tu,
                "dvt": dvt,
                "ma_quan_ly": ma_quan_ly,
                "ten_quan_ly": ten_quan_ly,
            }
        if ma_quan_ly and ten_quan_ly:
            manager_names.setdefault(ma_quan_ly, ten_quan_ly)
        if not ma_quan_ly:
            missing_manager_rows += 1

        key = (don_vi, kho_xuat, ma_hang, nam, thang)
        if key in fact_keys:
            duplicate_fact_keys += 1
        fact_keys.add(key)

        for header, value in zip(headers, values):
            if value in (None, ""):
                nulls[header] += 1

    wb.close()
    return {
        "headers": headers,
        "rows": rows,
        "units": units,
        "warehouses": warehouses,
        "codes": codes,
        "manager_names": manager_names,
        "missing_manager_rows": missing_manager_rows,
        "date_min": min(dates),
        "date_max": max(dates),
        "duplicate_fact_keys": duplicate_fact_keys,
        "nulls": nulls,
    }


def make_reconciliation_sheet(wb, usage, catalog, catalog_rows, catalog_nulls):
    ws = wb.create_sheet("02_DOI_CHIEU_2_FILE", 2)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "KẾT QUẢ ĐỐI CHIẾU HAI FILE PRODUCTION"
    style_title(ws["A1"], 18)
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Các số liệu dưới đây được tính trực tiếp từ hai file trong folder database. "
        "Workbook này đã nạp sẵn danh mục ghép và 141.623 dòng lịch sử chuẩn hóa."
    )
    ws["A2"].fill = fill(BLUE)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 38

    catalog_codes = set(catalog)
    usage_codes = set(usage["codes"])
    duplicate_codes = {code for code, rows in catalog.items() if len(rows) > 1}
    missing_manager_codes = {
        code for code, row in usage["codes"].items() if not row["ma_quan_ly"]
    }
    metrics = (
        ("Dòng lịch sử sử dụng hợp lệ", usage["rows"], "Đã nạp vào LICH_SU_XUAT_KHO"),
        ("Khoảng dữ liệu", f'{usage["date_min"]} → {usage["date_max"]}', "30 tháng"),
        ("Đơn vị có phát sinh", len(usage["units"]), "Cần danh sách đơn vị chính thức để biết khoa còn thiếu"),
        ("Kho xuất", len(usage["warehouses"]), "Giữ nguyên chi tiết, không gộp lúc nạp"),
        ("Mã hàng có sử dụng", len(usage_codes), "Nguồn từ file số lượng"),
        ("Mã danh mục production", len(catalog_codes), f"{catalog_rows} dòng nguồn"),
        ("Mã sử dụng có đặc tả", len(usage_codes & catalog_codes), f"{len(usage_codes & catalog_codes) / len(usage_codes):.2%} số mã đang dùng"),
        ("Mã sử dụng thiếu đặc tả", len(usage_codes - catalog_codes), "Phải bổ sung gói + tiêu chí kỹ thuật"),
        ("Mã danh mục chưa phát sinh 30 tháng", len(catalog_codes - usage_codes), "Vẫn giữ trong danh mục"),
        ("Mã sử dụng thiếu mã quản lý", len(missing_manager_codes), f'{usage["missing_manager_rows"]} dòng lịch sử'),
        ("Mã danh mục bị lặp/xung đột", len(duplicate_codes), "Không tự chọn một bản; bệnh viện phải xác nhận"),
        ("Khóa fact bị trùng", usage["duplicate_fact_keys"], "0 là đạt"),
    )
    row = 4
    ws.append(())
    for col, value in enumerate(("Chỉ số", "Kết quả", "Nhận định"), 1):
        ws.cell(row, col, value)
        style_header(ws.cell(row, col))
    for label, value, note in metrics:
        row += 1
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 3, note)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        if any(key in label for key in ("thiếu", "xung đột", "bị trùng")) and value:
            for col in range(1, 4):
                ws.cell(row, col).fill = fill(YELLOW if "bị trùng" not in label else RED)

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "ÁNH XẠ CỘT FILE PRODUCTION → WORKBOOK/DB").fill = fill(NAVY)
    ws.cell(row, 1).font = Font(bold=True, color=WHITE)
    row += 1
    headers = ("File nguồn", "Cột nguồn", "Cột đích", "Bảng/view đích", "Xử lý")
    for col, value in enumerate(headers, 1):
        ws.cell(row, col, value)
        style_header(ws.cell(row, col))
    mappings = (
        (USAGE_NAME, "Đơn vị", "don_vi", "usage_history_current", "Giữ nguyên; chuẩn hóa tên khoa sau đối soát"),
        (USAGE_NAME, "Kho xuất", "kho_xuat", "usage_history_current", "Giữ nguyên chi tiết"),
        (USAGE_NAME, "Mã quản lý", "ma_quan_ly", "nhom_ky_thuat/vat_tu", "Giữ text; 511 mã hàng đang trống"),
        (USAGE_NAME, "Tên quản lý", "ten_quan_ly", "nhom_ky_thuat", "Lấy theo mã quản lý"),
        (USAGE_NAME, "Mã hàng", "ma_hang", "vat_tu/usage_history_current", "Ép text"),
        (USAGE_NAME, "Tên vật tư", "ten_vat_tu", "vat_tu", "Nguồn tên cơ bản khi mã có sử dụng"),
        (USAGE_NAME, "ĐVT", "dvt", "vat_tu", "Cần bảng quy đổi nếu từng thay đổi"),
        (USAGE_NAME, "Ngày", "nam + thang", "usage_history_current", "Tách năm/tháng; giữ mức tháng"),
        (USAGE_NAME, "Số lượng", "so_luong", "usage_history_current", "Số không âm"),
        (CATALOG_NAME, "Gói", "goi", "vat_tu", "Lấy khi mã không bị lặp xung đột"),
        (CATALOG_NAME, "Tiêu chí kỹ thuật", "tieu_chi_ky_thuat", "vat_tu", "Lấy khi mã không bị lặp xung đột"),
        (CATALOG_NAME, "Tên thương mại", "ten_thuong_mai", "vat_tu", "Trống ở một số mã"),
        (CATALOG_NAME, "Ký mã hiệu", "ky_ma_hieu", "vat_tu", "Trống ở một số mã"),
        (CATALOG_NAME, "Hãng", "hang", "vat_tu", "Trống ở một số mã"),
        (CATALOG_NAME, "Nước sản xuất", "nuoc_san_xuat", "vat_tu", "Trống ở một số mã"),
    )
    for mapping in mappings:
        row += 1
        for col, value in enumerate(mapping, 1):
            ws.cell(row, col, value)

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "DỮ LIỆU HAI FILE CHƯA CÓ — CẦN ĐIỀN CÁC SHEET CÒN LẠI").fill = fill(NAVY)
    ws.cell(row, 1).font = Font(bold=True, color=WHITE)
    missing = (
        ("DON_VI_TAI_KHOAN", "Danh sách đơn vị chính thức, email, họ tên, role; không gửi mật khẩu."),
        ("DOT_DE_XUAT / DE_XUAT_CU / LY_DO_DE_XUAT", "Đợt và đề xuất lịch sử để hiển thị workflow cũ."),
        ("KY_THAU_SO_CHOT", "Kỳ thầu và số đã chốt; staging hiện 0 dòng."),
        ("GOI_THAU_TIMELINE / KET_QUA_THAU", "Toàn bộ gói, năm mốc, số đề xuất/trúng, mã rớt và lý do."),
        ("HOP_DONG / TON_KHO_HANG_VE", "Hợp đồng, giá, tồn kho, hàng đang về và lead time."),
        ("THIEU_HANG / XAC_NHAN_THANG", "Bằng chứng nhu cầu bị nén vì cấp hạn chế/hết hàng."),
        ("SU_KIEN_NHU_CAU", "Kỹ thuật mới, thay đổi phác đồ, tăng/giảm ca có định lượng."),
        ("HO_SO_CU_MANIFEST", "Danh sách file Word/Excel cũ; gửi file gốc riêng."),
        ("BIEU_MAU_CHINH_THUC", "Năm mẫu chính thức + file đã điền mẫu."),
        ("QUY_DOI_THAY_THE", "ĐVT/quy cách, đổi mã và thay thế tương đương lâm sàng."),
    )
    for sheet, detail in missing:
        row += 1
        ws.cell(row, 1, sheet).font = Font(bold=True, color=NAVY)
        ws.cell(row, 2, detail)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)

    for idx, width in enumerate((28, 28, 30, 26, 44, 18, 18, 18), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"


def make_issues_sheet(wb, usage, catalog, catalog_nulls):
    ws = wb.create_sheet("03_VAN_DE_DU_LIEU", 3)
    ws.sheet_view.showGridLines = False
    headers = (
        "loai_van_de",
        "ma_hang",
        "ma_quan_ly_hien_co",
        "ten_vat_tu_hien_co",
        "nguon",
        "chi_tiet",
        "hanh_dong_can_lam",
        "trang_thai_xu_ly",
    )
    ws.append(headers)
    for cell in ws[1]:
        style_header(cell)

    usage_codes = usage["codes"]
    usage_set = set(usage_codes)
    catalog_set = set(catalog)

    for code in sorted(usage_set - catalog_set):
        info = usage_codes[code]
        ws.append(
            (
                "THIEU_DAC_TA",
                code,
                info["ma_quan_ly"],
                info["ten_vat_tu"],
                USAGE_NAME,
                "Mã có sử dụng nhưng không có trong file thông tin VTYT.",
                "Bổ sung gói, tiêu chí kỹ thuật, tên thương mại, ký mã hiệu, hãng, nước sản xuất.",
                "cho_bo_sung",
            )
        )

    for code, info in sorted(usage_codes.items()):
        if not info["ma_quan_ly"]:
            ws.append(
                (
                    "THIEU_MA_QUAN_LY",
                    code,
                    "",
                    info["ten_vat_tu"],
                    USAGE_NAME,
                    "Mã hàng chưa được gán mã quản lý/nhóm kỹ thuật.",
                    "Xác nhận ma_quan_ly và ten_quan_ly; nếu chưa có nhóm ghi rõ lý do.",
                    "cho_bo_sung",
                )
            )

    for code, rows in sorted(catalog.items()):
        if len(rows) <= 1:
            continue
        for index, row in enumerate(rows, 1):
            detail = " | ".join(
                f"{key}={clean(row.get(key))}"
                for key in ("Gói", "Tên vật tư", "ĐVT", "Tên thương mại", "Ký mã hiệu", "Hãng", "Nước sản xuất")
            )
            ws.append(
                (
                    "XUNG_DOT_DANH_MUC",
                    code,
                    usage_codes.get(code, {}).get("ma_quan_ly", ""),
                    usage_codes.get(code, {}).get("ten_vat_tu", clean(row.get("Tên vật tư"))),
                    f"{CATALOG_NAME} bản {index}/{len(rows)}",
                    detail,
                    "Chọn một bản đúng hoặc xác nhận đây thực sự là hai mã khác nhau bị dùng chung ma_hang.",
                    "cho_xac_nhan",
                )
            )

    for code in sorted(catalog_set - usage_set):
        row = catalog[code][0]
        ws.append(
            (
                "CHUA_CO_LICH_SU_SU_DUNG",
                code,
                "",
                clean(row.get("Tên vật tư")),
                CATALOG_NAME,
                "Có trong danh mục nhưng không phát sinh trong 30 tháng của file sử dụng.",
                "Xác nhận còn hoạt động; không xóa nếu là mã dự phòng/mới.",
                "cho_xac_nhan",
            )
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[0].value == "XUNG_DOT_DANH_MUC":
            for cell in row:
                cell.fill = fill(RED)
        elif row[0].value in {"THIEU_DAC_TA", "THIEU_MA_QUAN_LY"}:
            for cell in row:
                cell.fill = fill(YELLOW)
    widths = (25, 16, 22, 38, 30, 70, 62, 18)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"


def fill_master(ws, usage, catalog):
    usage_codes = usage["codes"]
    all_codes = set(usage_codes) | set(catalog)
    for code in sorted(all_codes, key=lambda value: (not value.isdigit(), int(value) if value.isdigit() else value)):
        basic = usage_codes.get(code, {})
        catalog_rows = catalog.get(code, [])
        unique_catalog = catalog_rows[0] if len(catalog_rows) == 1 else {}
        ten_vat_tu = basic.get("ten_vat_tu") or clean(unique_catalog.get("Tên vật tư"))
        dvt = basic.get("dvt") or clean(unique_catalog.get("ĐVT"))
        ma_quan_ly = basic.get("ma_quan_ly", "")
        ten_quan_ly = basic.get("ten_quan_ly", "")
        if len(catalog_rows) > 1:
            source = f"{USAGE_NAME if basic else ''}; {CATALOG_NAME}: XUNG ĐỘT {len(catalog_rows)} dòng — xem 03_VAN_DE_DU_LIEU".strip("; ")
        elif basic and unique_catalog:
            source = f"{USAGE_NAME} + {CATALOG_NAME}"
        elif basic:
            source = f"{USAGE_NAME}; THIẾU ĐẶC TẢ"
        else:
            source = f"{CATALOG_NAME}; CHƯA CÓ LỊCH SỬ SỬ DỤNG"
        ws.append(
            (
                code,
                ten_vat_tu,
                dvt,
                ma_quan_ly,
                ten_quan_ly,
                clean(unique_catalog.get("Gói")),
                clean(unique_catalog.get("Tiêu chí kỹ thuật")),
                clean(unique_catalog.get("Tên thương mại")),
                clean(unique_catalog.get("Ký mã hiệu")),
                clean(unique_catalog.get("Hãng")),
                clean(unique_catalog.get("Nước sản xuất")),
                True if basic else None,
                source,
            )
        )


def fill_units(ws, usage):
    for unit in sorted(usage["units"]):
        ws.append(("", unit, "", "", "", "", "Có phát sinh trong file lịch sử; bổ sung mã đơn vị và tài khoản."))


def fill_defaults(wb):
    ws = wb["MOC_CAM_KET"]
    ws.append((6, 0.20, "Sau 6 tháng phải dùng tối thiểu 20%"))
    ws.append((12, 0.50, "Sau 12 tháng phải dùng tối thiểu 50%"))
    ws.append((18, 0.80, "Hết kỳ 18 tháng phải đạt cam kết 80%"))

    ws = wb["BIEU_MAU_CHINH_THUC"]
    defaults = (
        ("mau_chi_dinh_thau", "Đề xuất mua chỉ định thầu", "chi_dinh_thau", "word"),
        ("mau_cam_ket_sl", "Bản cam kết số lượng", "cam_ket_sl", "word"),
        ("mau_danh_muc_dvsd", "Danh mục đề xuất ĐVSD", "danh_muc_dvsd", "excel"),
        ("mau_de_nghi_mua", "Đề nghị mua thầu của PĐD", "de_nghi_mua", "word"),
        ("mau_tong_hop_thau", "Tổng hợp danh mục đi thầu", "tong_hop_thau", "excel"),
    )
    for code, name, dossier, file_type in defaults:
        ws.append((code, name, dossier, file_type, "", None, None, "", "", True, "Bổ sung file trống mới nhất và một file đã điền mẫu."))


def create(source_dir: Path, output: Path):
    usage_path = source_dir / USAGE_NAME
    catalog_path = source_dir / CATALOG_NAME
    for path in (usage_path, catalog_path):
        if not path.exists():
            raise FileNotFoundError(path)

    build(output)
    wb = load_workbook(output)
    guide = wb["00_HUONG_DAN"]
    guide.merge_cells("A3:H3")
    guide["A3"] = (
        f"ĐÃ ĐỐI CHIẾU PRODUCTION: {USAGE_NAME} + {CATALOG_NAME}. "
        "Hai sheet DM_VAT_TU và LICH_SU_XUAT_KHO đã được nạp sẵn; không cần nhập lại."
    )
    guide["A3"].fill = fill(GREEN)
    guide["A3"].font = Font(bold=True, color=TEXT)
    guide["A3"].alignment = Alignment(wrap_text=True)
    guide.row_dimensions[3].height = 34

    catalog_headers, catalog, catalog_rows, catalog_nulls = read_catalog(catalog_path)
    usage = scan_usage_and_fill(usage_path, wb["LICH_SU_XUAT_KHO"])
    make_reconciliation_sheet(wb, usage, catalog, catalog_rows, catalog_nulls)
    make_issues_sheet(wb, usage, catalog, catalog_nulls)
    fill_master(wb["DM_VAT_TU"], usage, catalog)
    fill_units(wb["DON_VI_TAI_KHOAN"], usage)
    fill_defaults(wb)

    wb["DM_VAT_TU"].auto_filter.ref = f"A1:M{wb['DM_VAT_TU'].max_row}"
    wb["LICH_SU_XUAT_KHO"].auto_filter.ref = f"A1:I{wb['LICH_SU_XUAT_KHO'].max_row}"
    wb["DON_VI_TAI_KHOAN"].auto_filter.ref = f"A1:G{wb['DON_VI_TAI_KHOAN'].max_row}"
    wb.save(output)
    wb.close()

    check = load_workbook(output, read_only=True, data_only=False)
    assert check["LICH_SU_XUAT_KHO"].max_row - 1 == usage["rows"]
    assert check["DM_VAT_TU"].max_row - 1 == len(set(usage["codes"]) | set(catalog))
    assert "02_DOI_CHIEU_2_FILE" in check.sheetnames
    assert "03_VAN_DE_DU_LIEU" in check.sheetnames
    check.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=Path("database"))
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("database/DU_LIEU_FULL_WORKFLOW_DA_DOI_CHIEU_PRODUCTION.xlsx"),
    )
    args = parser.parse_args()
    create(args.source_dir.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
