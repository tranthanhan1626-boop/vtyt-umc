#!/usr/bin/env python3
"""Nạp "Số đề xuất kỳ trước" từ `de_xuat.xlsx` vào `de_xuat_ky_truoc`.

Bảng dựng ở `sql/patch_zzzzzzzg_de_xuat_ky_truoc.sql`. Chỉ để XEM trên web
(QĐ 18/09/2026 mục p): số khoa gõ ban đầu của kỳ thầu trước, 18 tháng.

    cd backend
    set -a && . ./.env.local && set +a

    # Xem trước (mặc định) — đọc file + đọc DB để đối chiếu, KHÔNG ghi:
    .venv/bin/python scripts/nap_de_xuat_ky_truoc.py ../de_xuat.xlsx --kiem

    # Ghi thật (upsert theo khoa + mã hàng + số QĐ, một transaction):
    .venv/bin/python scripts/nap_de_xuat_ky_truoc.py ../de_xuat.xlsx \\
        --that-su-nap --xac-nhan-staging

    # Chỉ làm sạch file, không mở kết nối nào (cần danh sách khoa từ CSV):
    .venv/bin/python scripts/nap_de_xuat_ky_truoc.py ../de_xuat.xlsx \\
        --khong-db --danh-sach-khoa <file.csv có cột ten_khoa>

Làm sạch:
  * Số QĐ: chữ Ð (U+00D0) → Đ (U+0110).
  * Mã hàng: ép sang text, bỏ khoảng trắng và xuống dòng ở hai đầu.
  * Số lượng: bỏ khoảng trắng (kể cả \\xa0) rồi đọc thành Decimal; giữ số lẻ.
  * Tên khoa: khớp đúng tên trên DB thì giữ; khớp sau khi bỏ khác biệt HOA/
    thường + khoảng trắng và chỉ có MỘT ứng viên thì ghép; còn lại BỎ và báo.
  * Năm đơn vị chủ dự án chốt bỏ (18/09): xem `BO_THEO_CHOT`.
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

STAGING_REF = "ihgfafubwyxnbubmppbj"
TEN_NGUON_MAC_DINH = "de_xuat.xlsx"

# Chủ dự án chốt bỏ qua ngày 18/09/2026.
BO_THEO_CHOT = frozenset({
    "Cơ sở 2",
    "Cơ sở 3",
    "Khoa GMHS - Hồi tỉnh",
    "Khoa Tuyến vú",
    "Đơn vị Hình ảnh tim mạch",
})

COT = ("Đơn vị", "Quyết định", "Mã hàng", "Số lượng đề xuất")


# ─── Hàm làm sạch thuần tuý (test bằng pytest) ────────────────────────────

def chuan_hoa_qd(v) -> str | None:
    if v is None:
        return None
    s = unicodedata.normalize("NFC", str(v)).strip().replace("Ð", "Đ")
    return s or None


def chuan_hoa_ma(v) -> str | None:
    """Mã hàng về text. Số nguyên từ Excel (62717) và chuỗi ('62717\\n') cùng
    ra '62717'. Không làm mất số 0 đầu của mã dạng chuỗi."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else None
    s = str(v).replace(" ", " ").strip()
    return s or None


def chuan_hoa_so(v) -> Decimal | None:
    """Số lượng về Decimal. '5\\xa0' → 5. Không đọc được → None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(v))
    s = str(v).replace(" ", "").strip()
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _khoa_so(s: str) -> str:
    return " ".join(unicodedata.normalize("NFC", s).split()).casefold()


# Tên trong file không tự ghép được (nhiều ứng viên) — chủ dự án chỉ định.
GHEP_TAY = {
    "KHOA GMHS (PHÒNG MỔ 2A)": "Khoa GMHS - Phòng mổ",   # QĐ 18/09/2026
}


def ghep_ten_khoa(ten: str, ds_khoa) -> tuple[str | None, str]:
    """Trả (tên khoa trên DB | None, lý do).

    lý do ∈ {"khop", "ghep", "bo_theo_chot", "khong_khop", "nhieu_ung_vien"}.
    """
    ten = unicodedata.normalize("NFC", str(ten or "")).strip()
    if ten in BO_THEO_CHOT:
        return None, "bo_theo_chot"
    ds = [unicodedata.normalize("NFC", k) for k in ds_khoa]
    # Ghép tay do chủ dự án chốt 18/09/2026 — chỉ nhận khi tên đích có thật trên DB.
    dich = GHEP_TAY.get(ten)
    if dich is not None:
        return (dich, "ghep") if dich in ds else (None, "khong_khop")
    if ten in ds:
        return ten, "khop"
    ung_vien = sorted({k for k in ds if _khoa_so(k) == _khoa_so(ten)})
    if len(ung_vien) == 1:
        return ung_vien[0], "ghep"
    if len(ung_vien) > 1:
        return None, "nhieu_ung_vien"
    return None, "khong_khop"


@dataclass
class BaoCao:
    tong_dong: int = 0
    giu: int = 0
    bo: Counter = field(default_factory=Counter)          # lý do → số dòng
    bo_theo_don_vi: Counter = field(default_factory=Counter)
    ghep: dict = field(default_factory=dict)               # tên gốc → tên DB
    ghep_so_dong: Counter = field(default_factory=Counter)
    ma_co_khoang_trang: int = 0
    so_sua: list = field(default_factory=list)             # (dòng Excel, gốc, sau)
    so_le: int = 0
    trung: list = field(default_factory=list)              # (khoa, ma, qd, [dòng])
    qd: Counter = field(default_factory=Counter)


def lam_sach(dong_tho, ds_khoa) -> tuple[list[dict], BaoCao]:
    """`dong_tho`: iterable (so_dong_excel, don_vi, quyet_dinh, ma_hang, so_luong).

    Trả danh sách dòng sạch {khoa, ma_hang, so_quyet_dinh, so_luong} và báo cáo.
    Không gộp dòng trùng: có trùng thì ghi vào `bc.trung`, người gọi phải dừng.
    """
    bc = BaoCao()
    sach: list[dict] = []
    vi_tri: dict[tuple, list[int]] = defaultdict(list)
    for so_dong, don_vi, qd_goc, ma_goc, so_goc in dong_tho:
        if all(x is None or str(x).strip() == "" for x in (don_vi, qd_goc, ma_goc, so_goc)):
            continue
        bc.tong_dong += 1
        khoa, ly_do = ghep_ten_khoa(don_vi, ds_khoa)
        if khoa is None:
            bc.bo[ly_do] += 1
            bc.bo_theo_don_vi[str(don_vi).strip()] += 1
            continue
        if ly_do == "ghep":
            bc.ghep[str(don_vi).strip()] = khoa
            bc.ghep_so_dong[str(don_vi).strip()] += 1
        qd = chuan_hoa_qd(qd_goc)
        ma = chuan_hoa_ma(ma_goc)
        so = chuan_hoa_so(so_goc)
        if isinstance(ma_goc, str) and ma_goc != ma_goc.strip():
            bc.ma_co_khoang_trang += 1
        if isinstance(so_goc, str) and so is not None:
            bc.so_sua.append((so_dong, so_goc, so))
        if qd is None:
            bc.bo["thieu_qd"] += 1
            continue
        if ma is None:
            bc.bo["thieu_ma"] += 1
            continue
        if so is None:
            bc.bo["so_khong_doc_duoc"] += 1
            continue
        if so < 0:
            bc.bo["so_am"] += 1
            continue
        if so != so.to_integral_value():
            bc.so_le += 1
        k = (khoa, ma, qd)
        vi_tri[k].append(so_dong)
        bc.qd[qd] += 1
        sach.append({"khoa": khoa, "ma_hang": ma, "so_quyet_dinh": qd, "so_luong": so,
                     "_dong": so_dong})
        bc.giu += 1
    bc.trung = [(*k, v) for k, v in vi_tri.items() if len(v) > 1]
    return sach, bc


# ─── Đọc file / DB ────────────────────────────────────────────────────────

def doc_xlsx(duong_dan: Path):
    from openpyxl import load_workbook
    wb = load_workbook(duong_dan, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    dau = tuple(str(x).strip() if x is not None else "" for x in next(it))
    if dau[:4] != COT:
        raise SystemExit(f"Tiêu đề cột không đúng mẫu: {dau[:4]} (cần {COT})")
    for i, r in enumerate(it, start=2):
        r = tuple(r) + (None,) * 4
        yield (i, r[0], r[1], r[2], r[3])


def doc_ds_khoa_csv(duong_dan: Path) -> list[str]:
    with open(duong_dan, encoding="utf-8-sig", newline="") as f:
        return [r["ten_khoa"].strip() for r in csv.DictReader(f) if (r.get("ten_khoa") or "").strip()]


def doc_ds_khoa_db(cur) -> list[str]:
    # v_don_vi = hợp tên khoa ở lịch sử HIS, users, proposals, khoa_nhom_ky_thuat
    # (patch_a5). Chạy bằng kết nối postgres nên không vướng RLS.
    cur.execute("select don_vi from v_don_vi")
    return [r[0] for r in cur.fetchall()]


def co_bang(cur) -> bool:
    cur.execute("select to_regclass('public.de_xuat_ky_truoc') is not null")
    return bool(cur.fetchone()[0])


# ─── In báo cáo ───────────────────────────────────────────────────────────

def in_bao_cao(bc: BaoCao, sach: list[dict]) -> None:
    print("── Báo cáo làm sạch ──")
    print(f"Dòng dữ liệu trong file : {bc.tong_dong:,}")
    print(f"Giữ                     : {bc.giu:,}")
    print(f"Bỏ                      : {sum(bc.bo.values()):,}")
    for ly_do, n in bc.bo.most_common():
        print(f"    {ly_do:<20} {n:,}")
    if bc.bo_theo_don_vi:
        print("  Đơn vị bị bỏ:")
        for ten, n in sorted(bc.bo_theo_don_vi.items()):
            print(f"    {ten:<40} {n:>5,} dòng")
    if bc.ghep:
        print("  Tên đã ghép về khoa trên DB:")
        for goc, db in sorted(bc.ghep.items()):
            print(f"    {goc!r:<40} → {db!r}  ({bc.ghep_so_dong[goc]} dòng)")
    print(f"Mã có khoảng trắng/xuống dòng đã cắt: {bc.ma_co_khoang_trang}")
    for so_dong, goc, sau in bc.so_sua:
        print(f"Ô số dạng chữ đã sửa: dòng Excel {so_dong}: {goc!r} → {sau}")
    print(f"Số lẻ giữ nguyên        : {bc.so_le}")
    print(f"Khoa                    : {len({r['khoa'] for r in sach})}")
    print(f"Mã hàng khác nhau       : {len({r['ma_hang'] for r in sach}):,}")
    print(f"Cặp (khoa, mã)          : {len({(r['khoa'], r['ma_hang']) for r in sach}):,}")
    print("Số QĐ (đã chuẩn hoá Đ U+0110):")
    for qd, n in sorted(bc.qd.items()):
        print(f"    {qd:<20} {n:>5,}")
    print(f"Dòng trùng (khoa, mã, QĐ): {len(bc.trung)}")
    for t in bc.trung[:20]:
        print(f"    {t}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    che_do = ap.add_mutually_exclusive_group()
    che_do.add_argument("--kiem", action="store_true", help="(mặc định) chỉ đọc, không ghi")
    che_do.add_argument("--that-su-nap", action="store_true", help="upsert vào de_xuat_ky_truoc")
    ap.add_argument("--xac-nhan-staging", action="store_true")
    ap.add_argument("--khong-db", action="store_true",
                    help="không mở kết nối DB; cần --danh-sach-khoa")
    ap.add_argument("--danh-sach-khoa", type=Path,
                    help="CSV có cột ten_khoa, dùng thay DB khi --khong-db")
    ap.add_argument("--nguon", default=TEN_NGUON_MAC_DINH)
    args = ap.parse_args(argv)

    if not args.file.exists():
        print(f"Không thấy file: {args.file}")
        return 1
    dong_tho = list(doc_xlsx(args.file))

    # ── Chế độ không DB: chỉ làm sạch ──
    if args.khong_db:
        if args.that_su_nap:
            print("TỪ CHỐI: --that-su-nap cần DB, không đi cùng --khong-db.")
            return 2
        if not args.danh_sach_khoa:
            print("--khong-db cần --danh-sach-khoa <csv> để ghép tên khoa.")
            return 2
        sach, bc = lam_sach(dong_tho, doc_ds_khoa_csv(args.danh_sach_khoa))
        in_bao_cao(bc, sach)
        print("\n(--khong-db: không đối chiếu DB, không ghi gì.)")
        return 1 if bc.trung else 0

    dsn = os.environ.get("SUPABASE_STAGING_DB_URL", "").strip()
    if not dsn:
        print("Thiếu SUPABASE_STAGING_DB_URL (nạp backend/.env.local trước).")
        return 1
    if STAGING_REF not in dsn:
        print("TỪ CHỐI: SUPABASE_STAGING_DB_URL không trỏ project staging đã định danh.")
        return 2
    if args.that_su_nap and not args.xac_nhan_staging:
        print("Thiếu --xac-nhan-staging; chưa nạp.")
        return 2

    import psycopg

    with psycopg.connect(dsn) as cn:
        if not args.that_su_nap:
            cn.read_only = True
        with cn.cursor() as cur:
            ds_khoa = doc_ds_khoa_db(cur)
            sach, bc = lam_sach(dong_tho, ds_khoa)
            in_bao_cao(bc, sach)

            print("\n── Đối chiếu DB ──")
            ds_ma = sorted({r["ma_hang"] for r in sach})
            cur.execute("select ma_hang from vat_tu where ma_hang = any(%s)", (ds_ma,))
            co_ma = {r[0] for r in cur.fetchall()}
            thieu = [m for m in ds_ma if m not in co_ma]
            print(f"Mã hàng có trong vat_tu : {len(co_ma):,}/{len(ds_ma):,}")
            if thieu:
                print(f"Mã KHÔNG có trong vat_tu: {len(thieu)} (vẫn nạp; web chỉ hiện ở mã đang có) "
                      f"— vd {thieu[:10]}")

            bang = co_bang(cur)
            moi = doi = giong = 0
            if not bang:
                print("Bảng de_xuat_ky_truoc CHƯA có — chạy patch_zzzzzzzg trước.")
            else:
                cur.execute("select khoa, ma_hang, so_quyet_dinh, so_luong from de_xuat_ky_truoc")
                cu = {(k, m, q): s for k, m, q, s in cur.fetchall()}
                for r in sach:
                    k = (r["khoa"], r["ma_hang"], r["so_quyet_dinh"])
                    if k not in cu:
                        moi += 1
                    elif Decimal(cu[k]) != r["so_luong"]:
                        doi += 1
                    else:
                        giong += 1
                chi_db = len(set(cu) - {(r["khoa"], r["ma_hang"], r["so_quyet_dinh"]) for r in sach})
                print(f"Trên DB đang có {len(cu):,} dòng · sẽ thêm {moi:,} · sẽ đổi số {doi:,} · "
                      f"giữ nguyên {giong:,} · có trên DB mà không có trong file {chi_db:,} (không xoá)")

            if bc.trung:
                print("\n⛔ Có dòng trùng (khoa, mã, QĐ) — không nạp.")
                return 1
            if not args.that_su_nap:
                print("\n(--kiem: không ghi gì. Ghi thật: --that-su-nap --xac-nhan-staging)")
                return 0
            if not bang:
                return 1

            cur.executemany(
                """insert into de_xuat_ky_truoc (khoa, ma_hang, so_quyet_dinh, so_luong, nguon, nap_luc)
                   values (%s, %s, %s, %s, %s, now())
                   on conflict (khoa, ma_hang, so_quyet_dinh)
                   do update set so_luong = excluded.so_luong,
                                 nguon    = excluded.nguon,
                                 nap_luc  = excluded.nap_luc""",
                [(r["khoa"], r["ma_hang"], r["so_quyet_dinh"], r["so_luong"], args.nguon) for r in sach],
            )
            cur.execute("select count(*) from de_xuat_ky_truoc")
            print(f"\n✅ Đã upsert {len(sach):,} dòng. Bảng nay có {cur.fetchone()[0]:,} dòng.")
        # thoát `with psycopg.connect` không lỗi → commit một lần.
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
