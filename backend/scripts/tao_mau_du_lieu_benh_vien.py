#!/usr/bin/env python3
"""Tạo workbook bàn giao dữ liệu bệnh viện cho toàn bộ workflow VTYT.

Workbook là biểu mẫu đầu vào, không phải bản dump database. Các ID, JSON audit
và timestamp do hệ thống tự sinh được thay bằng mã nguồn dễ đối soát.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


@dataclass(frozen=True)
class Field:
    name: str
    level: str
    dtype: str
    description: str
    target: str
    allowed: str = ""
    example: str = ""


@dataclass(frozen=True)
class SheetSpec:
    name: str
    title: str
    purpose: str
    target: str
    phase: str
    fields: tuple[Field, ...]


def f(
    name: str,
    level: str,
    dtype: str,
    description: str,
    target: str,
    allowed: str = "",
    example: str = "",
) -> Field:
    return Field(name, level, dtype, description, target, allowed, example)


LOAI_MUA = "dau_thau_rong_rai | mua_sam_bo_sung | chi_dinh_thau"
TRANG_THAI_DE_XUAT = "de_xuat | xet_duyet | hoan_thanh | tu_choi"
BOOL = "TRUE | FALSE"
DATE = "YYYY-MM-DD"
DATETIME = "YYYY-MM-DD HH:MM:SS+07:00"


SHEETS: tuple[SheetSpec, ...] = (
    SheetSpec(
        "DM_VAT_TU",
        "Danh mục vật tư và nhóm kỹ thuật",
        "Một dòng cho một mã hàng. Dùng để tạo nhom_ky_thuat và vat_tu; mã phải giữ dạng text.",
        "nhom_ky_thuat + vat_tu",
        "Tối thiểu",
        (
            f("ma_hang", "Bắt buộc", "Text", "Mã SKU/mã hàng duy nhất trong HIS; giữ số 0 đầu.", "vat_tu.ma_hang", example="0066114"),
            f("ten_vat_tu", "Bắt buộc", "Text", "Tên chuẩn của mã hàng tại bệnh viện.", "vat_tu.ten_vat_tu", example="Tăm bông vô khuẩn"),
            f("dvt", "Bắt buộc", "Text", "Đơn vị tính chuẩn dùng để đề xuất và so sánh lịch sử.", "vat_tu.dvt", example="Que"),
            f("ma_quan_ly", "Khuyến nghị", "Text", "Mã quản lý/nhóm kỹ thuật; một mã hàng chỉ thuộc tối đa một nhóm.", "vat_tu.ma_quan_ly -> nhom_ky_thuat.ma_quan_ly", example="N01.01.020.01"),
            f("ten_quan_ly", "Có điều kiện", "Text", "Bắt buộc khi có ma_quan_ly; phải thống nhất cho mọi dòng cùng mã quản lý.", "nhom_ky_thuat.ten_quan_ly", example="Tăm bông"),
            f("goi", "Khuyến nghị", "Text", "Nhãn gói chuyên môn hiện hành của vật tư.", "vat_tu.goi", example="Dùng chung"),
            f("tieu_chi_ky_thuat", "Khuyến nghị", "Text dài", "Tiêu chí kỹ thuật dùng để sinh hồ sơ Word/Excel.", "vat_tu.tieu_chi_ky_thuat"),
            f("ten_thuong_mai", "Khuyến nghị", "Text", "Tên thương mại nếu có.", "vat_tu.ten_thuong_mai"),
            f("ky_ma_hieu", "Khuyến nghị", "Text", "Ký/mã hiệu, catalogue hoặc model.", "vat_tu.ky_ma_hieu"),
            f("hang", "Khuyến nghị", "Text", "Hãng sản xuất.", "vat_tu.hang"),
            f("nuoc_san_xuat", "Khuyến nghị", "Text", "Nước sản xuất.", "vat_tu.nuoc_san_xuat"),
            f("dang_hoat_dong", "Khuyến nghị", "Boolean", "TRUE nếu còn dùng; giúp giữ mã cũ mà không xóa lịch sử.", "Cần bổ sung khi import", BOOL, "TRUE"),
            f("nguon_danh_muc", "Khuyến nghị", "Text", "Tên hệ thống/file nguồn và ngày chốt danh mục.", "Metadata import", example="HIS_2026-07-31"),
        ),
    ),
    SheetSpec(
        "LICH_SU_XUAT_KHO",
        "Lịch sử xuất/sử dụng theo tháng",
        "Một dòng theo đơn vị × kho xuất × mã hàng × năm × tháng. Nên cung cấp ít nhất 24–36 tháng.",
        "import_batches + usage_history_current",
        "Tối thiểu",
        (
            f("don_vi", "Bắt buộc", "Text", "Tên khoa/phòng sử dụng; phải giống tuyệt đối ở mọi sheet.", "usage_history_current.don_vi"),
            f("kho_xuat", "Bắt buộc", "Text", "Kho/máy/tủ xuất. Nếu không rõ để chuỗi rỗng, không dùng NULL.", "usage_history_current.kho_xuat"),
            f("ma_hang", "Bắt buộc", "Text", "Mã hàng phải tồn tại trong DM_VAT_TU.", "usage_history_current.ma_hang"),
            f("nam", "Bắt buộc", "Integer", "Năm ghi nhận sử dụng.", "usage_history_current.nam", "2000–2100", "2025"),
            f("thang", "Bắt buộc", "Integer", "Tháng ghi nhận sử dụng.", "usage_history_current.thang", "1–12", "7"),
            f("so_luong", "Bắt buộc", "Number", "Số lượng xuất/sử dụng trong tháng, không âm.", "usage_history_current.so_luong", ">= 0", "1250"),
            f("source_filename", "Bắt buộc", "Text", "Tên file nguồn để tạo import batch và truy vết.", "import_batches.source_filename"),
            f("ngay_chot_du_lieu", "Khuyến nghị", "Date", "Ngày dữ liệu được trích xuất/chốt từ HIS.", "import_batches.notes", DATE),
            f("ghi_chu_nguon", "Tùy chọn", "Text", "Cảnh báo Power BI cắt dòng, dữ liệu thiếu hoặc cách quy đổi.", "import_batches.warnings/notes"),
        ),
    ),
    SheetSpec(
        "DON_VI_TAI_KHOAN",
        "Danh sách đơn vị và tài khoản",
        "Không cung cấp mật khẩu. Tên đơn vị là khóa liên kết mềm quan trọng nhất giữa các sheet.",
        "users; danh mục đơn vị chuẩn hóa",
        "Tối thiểu",
        (
            f("ma_don_vi", "Khuyến nghị", "Text", "Mã khoa/phòng ổn định của bệnh viện; DB hiện chưa có cột riêng.", "Cần bảng ánh xạ khi import"),
            f("ten_don_vi", "Bắt buộc", "Text", "Tên khoa/phòng chuẩn; ánh xạ vào users.khoa và các cột don_vi.", "users.khoa / *.don_vi"),
            f("email", "Bắt buộc", "Email", "Email đăng nhập; không gửi mật khẩu.", "users.email", example="ten@umc.edu.vn"),
            f("ho_ten", "Bắt buộc", "Text", "Họ tên hiển thị và ghi dấu vết.", "users.ho_ten"),
            f("role", "Bắt buộc", "Enum", "Vai trò tài khoản.", "users.role", "dvsd | dieu_duong | admin", "dvsd"),
            f("dang_hoat_dong", "Bắt buộc", "Boolean", "TRUE nếu tài khoản/đơn vị còn được sử dụng.", "Metadata import", BOOL, "TRUE"),
            f("ghi_chu", "Tùy chọn", "Text", "Thông tin đối soát; không ghi bí mật đăng nhập.", "Metadata import"),
        ),
    ),
    SheetSpec(
        "DOT_DE_XUAT",
        "Đợt đề xuất",
        "Một dòng cho một đợt/gói do Phòng Điều dưỡng mở.",
        "dot_de_xuat",
        "Tối thiểu",
        (
            f("ma_dot_nguon", "Bắt buộc", "Text", "Mã đối soát do bệnh viện cung cấp; importer đổi thành dot_id.", "Khóa ánh xạ import", example="DX-2026-RR-01"),
            f("loai_mua_sam", "Bắt buộc", "Enum", "Loại mua sắm kỹ thuật của hệ thống.", "dot_de_xuat.loai_mua_sam", LOAI_MUA),
            f("ten", "Bắt buộc", "Text", "Tên đợt hiển thị trên web.", "dot_de_xuat.ten"),
            f("nam", "Bắt buộc", "Integer", "Năm của đợt.", "dot_de_xuat.nam", "2000–2100"),
            f("thang_moc", "Có điều kiện", "Integer", "Gói bổ sung dùng 1, 5 hoặc 9; gói khác có thể để trống.", "dot_de_xuat.thang_moc", "1 | 5 | 9 | trống"),
            f("trang_thai", "Bắt buộc", "Enum", "Đợt đang mở hay đã đóng.", "dot_de_xuat.trang_thai", "mo | dong"),
            f("ngay_mo", "Khuyến nghị", "Datetime", "Thời điểm mở đợt.", "dot_de_xuat.ngay_mo", DATETIME),
            f("ngay_dong", "Có điều kiện", "Datetime", "Thời điểm đóng; để trống nếu còn mở.", "dot_de_xuat.ngay_dong", DATETIME),
            f("ghi_chu", "Tùy chọn", "Text", "Hướng dẫn hoặc phạm vi đợt.", "dot_de_xuat.ghi_chu"),
            f("created_by", "Khuyến nghị", "Email", "Email người tạo đợt.", "dot_de_xuat.created_by"),
        ),
    ),
    SheetSpec(
        "DE_XUAT_CU",
        "Dòng đề xuất lịch sử",
        "Một dòng cho một phiên bản đề xuất mã hàng. Không tự tạo ID database.",
        "proposals",
        "Tối thiểu",
        (
            f("ma_de_xuat_nguon", "Bắt buộc", "Text", "Mã dòng nguồn để nối sang LY_DO_DE_XUAT.", "Khóa ánh xạ import"),
            f("ma_dot_nguon", "Bắt buộc", "Text", "Khớp DOT_DE_XUAT.ma_dot_nguon.", "proposals.dot_id qua ánh xạ"),
            f("ma_hang", "Bắt buộc", "Text", "Mã hàng đề xuất.", "proposals.ma_hang"),
            f("don_vi", "Bắt buộc", "Text", "Khoa/phòng đề xuất.", "proposals.don_vi"),
            f("nam_de_xuat", "Bắt buộc", "Integer", "Năm tài chính của nhu cầu.", "proposals.nam_de_xuat"),
            f("version", "Bắt buộc", "Integer", "Phiên bản tăng dần của cùng mã hàng/khoa/năm.", "proposals.version", ">= 1", "1"),
            f("is_current", "Bắt buộc", "Boolean", "Chỉ một version hiện hành được TRUE.", "proposals.is_current", BOOL),
            f("so_luong", "Bắt buộc", "Number", "Tổng số lượng đề xuất cho kỳ.", "proposals.so_luong", ">= 0"),
            f("so_thang_du_kien", "Bắt buộc", "Integer", "Số tháng tính cả tháng đầu và cuối; importer sẽ kiểm lại từ bốn mốc.", "proposals.so_thang_du_kien", "1–60"),
            f("loai_mua_sam", "Bắt buộc", "Enum", "Phải khớp loại của đợt.", "proposals.loai_mua_sam", LOAI_MUA),
            f("tu_thang", "Bắt buộc", "Integer", "Tháng bắt đầu sử dụng.", "proposals.tu_thang", "1–12"),
            f("tu_nam", "Bắt buộc", "Integer", "Năm bắt đầu sử dụng.", "proposals.tu_nam", "2000–2100"),
            f("den_thang", "Bắt buộc", "Integer", "Tháng kết thúc sử dụng.", "proposals.den_thang", "1–12"),
            f("den_nam", "Bắt buộc", "Integer", "Năm kết thúc; không trước kỳ bắt đầu.", "proposals.den_nam", "2000–2100"),
            f("nhom_de_xuat_nguon", "Khuyến nghị", "Text/UUID", "Các dòng cùng giỏ dùng chung một mã nhóm.", "proposals.nhom_de_xuat"),
            f("goi", "Khuyến nghị", "Text", "Nhãn gói chuyên môn snapshot tại thời điểm gửi.", "proposals.goi"),
            f("created_by", "Bắt buộc", "Email", "Email người lập, giữ dạng snapshot.", "proposals.created_by"),
            f("created_by_ho_ten", "Khuyến nghị", "Text", "Họ tên người lập tại thời điểm gửi.", "proposals.created_by_ho_ten"),
            f("created_at", "Bắt buộc", "Datetime", "Thời điểm tạo đề xuất.", "proposals.created_at", DATETIME),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái workflow cũ.", "proposals.trang_thai", TRANG_THAI_DE_XUAT),
            f("ly_do_tra_lai", "Có điều kiện", "Text", "Bắt buộc nếu trạng thái tu_choi.", "proposals.ly_do_tra_lai"),
            f("da_rut", "Bắt buộc", "Boolean", "TRUE nếu nhóm đề xuất đã được rút.", "proposals.da_rut", BOOL),
            f("rut_luc", "Có điều kiện", "Datetime", "Bắt buộc khi da_rut=TRUE.", "proposals.rut_luc", DATETIME),
            f("rut_boi", "Có điều kiện", "Email", "Người rút đề xuất.", "proposals.rut_boi"),
            f("ly_do_rut", "Có điều kiện", "Text", "Lý do rút.", "proposals.ly_do_rut"),
        ),
    ),
    SheetSpec(
        "LY_DO_DE_XUAT",
        "Lý do và giải trình đề xuất",
        "Nối với DE_XUAT_CU bằng ma_de_xuat_nguon. Ghi chú thêm là tùy chọn.",
        "proposal_reasons",
        "Tối thiểu",
        (
            f("ma_de_xuat_nguon", "Bắt buộc", "Text", "Khớp mã dòng trong DE_XUAT_CU.", "proposal_reasons.proposal_id qua ánh xạ"),
            f("loai_ly_do", "Bắt buộc", "Enum", "Nhóm lý do giải trình.", "proposal_reasons.loai_ly_do", "theo_lich_su | ky_thuat_moi | thay_doi_phac_do | khac"),
            f("ten_ky_thuat_moi", "Có điều kiện", "Text", "Bắt buộc khi loai_ly_do=ky_thuat_moi.", "proposal_reasons.ten_ky_thuat_moi"),
            f("uoc_ca_thang", "Có điều kiện", "Number", "Ước số ca/tháng của kỹ thuật mới.", "proposal_reasons.uoc_ca_thang", ">= 0"),
            f("ghi_chu", "Tùy chọn", "Text", "Ghi chú bổ sung; không bắt buộc.", "proposal_reasons.ghi_chu"),
        ),
    ),
    SheetSpec(
        "KY_THAU_SO_CHOT",
        "Kỳ thầu và số lượng đã chốt",
        "Lặp thông tin kỳ cho từng mã/khoa. Đây là dữ liệu staging đang thiếu hoàn toàn.",
        "ky_thau + so_luong_ky",
        "Bắt buộc để theo dõi sau thầu",
        (
            f("ma_ky_nguon", "Bắt buộc", "Text", "Mã đối soát kỳ thầu.", "Khóa ánh xạ import"),
            f("ten_ky", "Bắt buộc", "Text", "Tên kỳ thầu duy nhất.", "ky_thau.ten"),
            f("tu_thang", "Bắt buộc", "Integer", "Tháng đầu kỳ.", "ky_thau.tu_thang", "1–12"),
            f("tu_nam", "Bắt buộc", "Integer", "Năm đầu kỳ.", "ky_thau.tu_nam"),
            f("den_thang", "Bắt buộc", "Integer", "Tháng cuối kỳ.", "ky_thau.den_thang", "1–12"),
            f("den_nam", "Bắt buộc", "Integer", "Năm cuối kỳ.", "ky_thau.den_nam"),
            f("ngay_chot_so", "Khuyến nghị", "Date", "Ngày PĐD chốt sổ.", "ky_thau.ngay_chot_so", DATE),
            f("ngay_lo_dau_du_kien", "Khuyến nghị", "Date", "Ngày dự kiến lô đầu về.", "ky_thau.ngay_lo_dau_du_kien", DATE),
            f("ghi_chu_ky", "Tùy chọn", "Text", "Ghi chú chung của kỳ.", "ky_thau.ghi_chu"),
            f("ma_quan_ly", "Khuyến nghị", "Text", "Mã quản lý của dòng số chốt.", "so_luong_ky.ma_quan_ly"),
            f("ma_hang", "Bắt buộc", "Text", "Mã hàng được chốt.", "so_luong_ky.ma_hang"),
            f("don_vi", "Bắt buộc", "Text", "Khoa sở hữu số lượng.", "so_luong_ky.don_vi"),
            f("sl_khoa_de_xuat", "Khuyến nghị", "Number", "Số khoa đã đề xuất trước điều chỉnh.", "so_luong_ky.sl_khoa_de_xuat", ">= 0"),
            f("sl_chot", "Bắt buộc", "Number", "Số lượng cuối cùng được chốt.", "so_luong_ky.sl_chot", ">= 0"),
            f("cach_tinh", "Bắt buộc", "Text/Enum", "Phương pháp dùng để chốt, phục vụ backtest.", "so_luong_ky.cach_tinh", "trung_binh_thang | nam_cao_nhat_cong_20 | khoa_tu_de_xuat | khac"),
            f("ly_do", "Có điều kiện", "Text", "Giải thích khi điều chỉnh số khoa đề xuất hoặc chọn phương pháp khác.", "so_luong_ky.ly_do"),
            f("nguoi_duyet", "Khuyến nghị", "Email/Text", "Người duyệt/chốt số.", "so_luong_ky.nguoi_duyet"),
            f("ngay_chot", "Khuyến nghị", "Date", "Ngày chốt dòng.", "so_luong_ky.ngay_chot", DATE),
        ),
    ),
    SheetSpec(
        "GOI_THAU_TIMELINE",
        "Gói thầu và timeline 5 mốc",
        "Lặp thông tin gói cho mỗi mốc. Mốc chuẩn: chào giá, mở thầu, đánh giá, ký hợp đồng, hàng về đợt đầu.",
        "goi_thau_tien_do + goi_thau_moc",
        "Bắt buộc để chạy tab Tiến độ",
        (
            f("ma_goi_nguon", "Bắt buộc", "Text", "Mã gói đối soát xuyên suốt workbook.", "Khóa ánh xạ import"),
            f("ten_goi", "Bắt buộc", "Text", "Tên gói hiển thị.", "goi_thau_tien_do.ten_goi"),
            f("loai_mua_sam", "Bắt buộc", "Enum", "Loại mua sắm.", "goi_thau_tien_do.loai_mua_sam", LOAI_MUA),
            f("nam", "Bắt buộc", "Integer", "Năm gói.", "goi_thau_tien_do.nam"),
            f("ghi_chu_goi", "Tùy chọn", "Text", "Ghi chú chung.", "goi_thau_tien_do.ghi_chu"),
            f("created_by", "Khuyến nghị", "Email", "Người tạo gói.", "goi_thau_tien_do.created_by"),
            f("ma_moc", "Bắt buộc", "Enum", "Mã mốc chuẩn.", "goi_thau_moc.ma_moc", "chao_gia | mo_thau | danh_gia | ky_hop_dong | hang_ve_dot_dau"),
            f("so_thu_tu", "Bắt buộc", "Integer", "Thứ tự 1–5 tương ứng mã mốc.", "goi_thau_moc.so_thu_tu", "1–5"),
            f("trang_thai_moc", "Bắt buộc", "Enum", "Trạng thái xử lý mốc.", "goi_thau_moc.trang_thai", "chua_bat_dau | dang_lam | hoan_thanh"),
            f("ngay", "Có điều kiện", "Date", "Ngày dự kiến hoặc thực tế của mốc; bắt buộc khi hoàn thành.", "goi_thau_moc.ngay", DATE),
            f("ghi_chu_moc", "Tùy chọn", "Text", "Vướng mắc hoặc kết quả tại mốc.", "goi_thau_moc.ghi_chu"),
            f("cap_nhat_boi", "Khuyến nghị", "Email", "Người cập nhật cuối.", "goi_thau_moc.cap_nhat_boi"),
            f("cap_nhat_luc", "Khuyến nghị", "Datetime", "Thời gian cập nhật cuối.", "goi_thau_moc.cap_nhat_luc", DATETIME),
        ),
    ),
    SheetSpec(
        "KET_QUA_THAU",
        "Kết quả thầu theo từng mã và khoa",
        "Một dòng cho gói × mã hàng × khoa. Kết quả không trúng bắt buộc có lý do và mốc rớt.",
        "goi_thau_ket_qua_ma",
        "Bắt buộc để thông báo rớt và tiến độ sử dụng",
        (
            f("ma_goi_nguon", "Bắt buộc", "Text", "Khớp GOI_THAU_TIMELINE.ma_goi_nguon.", "goi_thau_ket_qua_ma.goi_id qua ánh xạ"),
            f("ma_hang", "Bắt buộc", "Text", "Mã hàng có kết quả.", "goi_thau_ket_qua_ma.ma_hang"),
            f("don_vi", "Bắt buộc", "Text", "Khoa nhận thông báo kết quả.", "goi_thau_ket_qua_ma.don_vi"),
            f("ket_qua", "Bắt buộc", "Enum", "Kết quả hiện tại.", "goi_thau_ket_qua_ma.ket_qua", "cho_ket_qua | trung_thau | khong_trung"),
            f("ly_do_khong_trung", "Có điều kiện", "Text", "Bắt buộc khi ket_qua=khong_trung.", "goi_thau_ket_qua_ma.ly_do_khong_trung"),
            f("ma_moc_rot", "Có điều kiện", "Enum", "Bắt buộc khi không trúng; dùng cùng mã mốc timeline.", "goi_thau_ket_qua_ma.ma_moc_rot", "chao_gia | mo_thau | danh_gia | ky_hop_dong | hang_ve_dot_dau"),
            f("so_luong_de_xuat", "Bắt buộc", "Number", "Số lượng của khoa đi vào gói.", "goi_thau_ket_qua_ma.so_luong_de_xuat", ">= 0"),
            f("so_luong_trung", "Có điều kiện", "Number", "NULL khi chưa có kết quả; 0 nếu rớt; không vượt số đề xuất.", "goi_thau_ket_qua_ma.so_luong_trung", "0..so_luong_de_xuat"),
            f("khoa_da_xem", "Bắt buộc", "Boolean", "FALSE để web hiện thông báo cần xác nhận.", "goi_thau_ket_qua_ma.khoa_da_xem", BOOL, "FALSE"),
            f("cap_nhat_boi", "Khuyến nghị", "Email", "Người cập nhật kết quả.", "goi_thau_ket_qua_ma.cap_nhat_boi"),
            f("cap_nhat_luc", "Khuyến nghị", "Datetime", "Thời gian cập nhật.", "goi_thau_ket_qua_ma.cap_nhat_luc", DATETIME),
        ),
    ),
    SheetSpec(
        "HOP_DONG",
        "Hợp đồng ở cấp gói",
        "Thông tin hợp đồng hiện hỗ trợ trong DB. Chi tiết giao hàng theo mã nằm ở sheet TON_KHO_HANG_VE.",
        "hop_dong",
        "Bắt buộc để đủ sổ kỳ thầu",
        (
            f("ma_ky_nguon", "Bắt buộc", "Text", "Khớp KY_THAU_SO_CHOT.ma_ky_nguon.", "hop_dong.ky_thau_id qua ánh xạ"),
            f("ten_goi", "Bắt buộc", "Text", "Tên gói trong hợp đồng.", "hop_dong.ten_goi"),
            f("so_hop_dong", "Khuyến nghị", "Text", "Số hợp đồng; DB hiện cần mở rộng cột riêng.", "Cần bổ sung schema"),
            f("tran_hop_dong", "Khuyến nghị", "Number", "Giá trị trần hợp đồng.", "hop_dong.tran_hop_dong", ">= 0"),
            f("tuy_chon_30", "Bắt buộc", "Boolean", "Có điều khoản tùy chọn mua thêm 30% hay không.", "hop_dong.tuy_chon_30", BOOL),
            f("lich_giao_du_kien", "Khuyến nghị", "Text", "Mô tả lịch giao tổng quát theo hợp đồng.", "hop_dong.lich_giao_du_kien"),
            f("nha_cung_cap", "Khuyến nghị", "Text", "Nhà cung cấp trúng thầu; cần bổ sung schema để phân tích.", "Cần bổ sung schema"),
            f("ngay_ky", "Khuyến nghị", "Date", "Ngày ký hợp đồng.", "Cần bổ sung schema", DATE),
            f("ngay_het_han", "Khuyến nghị", "Date", "Ngày hết hiệu lực.", "Cần bổ sung schema", DATE),
            f("ghi_chu", "Tùy chọn", "Text", "Điều khoản hoặc ngoại lệ cần lưu ý.", "Metadata import"),
        ),
    ),
    SheetSpec(
        "THIEU_HANG",
        "Sự kiện thiếu hàng",
        "Một dòng cho một lần khoa yêu cầu/cảnh báo thiếu hàng.",
        "su_kien_thieu_hang",
        "Cần để hiệu chỉnh nhu cầu bị nén",
        (
            f("ma_su_kien_nguon", "Khuyến nghị", "Text", "Mã đối soát từ hệ thống nguồn.", "Khóa ánh xạ import"),
            f("don_vi", "Bắt buộc", "Text", "Khoa báo thiếu.", "su_kien_thieu_hang.don_vi"),
            f("ma_hang", "Có điều kiện", "Text", "Mã hàng; có thể trống nếu vật tư chưa có mã.", "su_kien_thieu_hang.ma_hang"),
            f("ten_vat_tu_tu_do", "Có điều kiện", "Text", "Bắt buộc nếu không có ma_hang.", "su_kien_thieu_hang.ten_vat_tu_tu_do"),
            f("ngay_bao", "Bắt buộc", "Date", "Ngày phát sinh/báo thiếu.", "su_kien_thieu_hang.ngay_bao", DATE),
            f("tinh_trang", "Bắt buộc", "Enum", "Mức đáp ứng.", "su_kien_thieu_hang.tinh_trang", "du_hang | cap_han_che | het_hang"),
            f("sl_yeu_cau", "Khuyến nghị", "Number", "Số khoa thực sự cần.", "su_kien_thieu_hang.sl_yeu_cau", ">= 0"),
            f("sl_duoc_cap", "Khuyến nghị", "Number", "Số kho thực cấp; cùng ĐVT với mã hàng.", "su_kien_thieu_hang.sl_duoc_cap", ">= 0"),
            f("co_hoan_ca", "Bắt buộc", "Boolean", "Có hoãn/hủy ca do thiếu hàng.", "su_kien_thieu_hang.co_hoan_ca", BOOL),
            f("so_ca_hoan", "Có điều kiện", "Integer", "Số ca bị ảnh hưởng; cần khi co_hoan_ca=TRUE.", "su_kien_thieu_hang.so_ca_hoan", ">= 0"),
            f("ma_thay_the", "Tùy chọn", "Text", "Mã hàng đã dùng thay thế.", "su_kien_thieu_hang.ma_thay_the"),
            f("ma_ly_do", "Khuyến nghị", "Text", "Mã A/B/C/D trong MA_LY_DO.", "su_kien_thieu_hang.ma_ly_do"),
            f("ghi_chu", "Tùy chọn", "Text", "Diễn giải tình huống.", "su_kien_thieu_hang.ghi_chu"),
            f("nguoi_bao", "Bắt buộc", "Email/Text", "Người báo sự kiện.", "su_kien_thieu_hang.nguoi_bao"),
            f("trang_thai_xu_ly", "Bắt buộc", "Enum", "Trạng thái PĐD xử lý.", "su_kien_thieu_hang.trang_thai_xu_ly", "moi_bao | da_xem | dang_xu_ly | da_xu_ly"),
            f("phan_hoi_pdd", "Tùy chọn", "Text", "Phản hồi của Phòng Điều dưỡng.", "su_kien_thieu_hang.phan_hoi_pdd"),
            f("nguoi_xac_nhan", "Tùy chọn", "Email/Text", "Người xác nhận/xử lý.", "su_kien_thieu_hang.nguoi_xac_nhan"),
            f("ngay_xac_nhan", "Tùy chọn", "Datetime", "Thời gian xác nhận.", "su_kien_thieu_hang.ngay_xac_nhan", DATETIME),
            f("an_khoi_bao_cao", "Bắt buộc", "Boolean", "Ẩn khỏi báo cáo nhưng không xóa dấu vết.", "su_kien_thieu_hang.an_khoi_bao_cao", BOOL, "FALSE"),
            f("ly_do_an", "Có điều kiện", "Text", "Bắt buộc nếu an_khoi_bao_cao=TRUE.", "su_kien_thieu_hang.ly_do_an"),
            f("nguoi_an", "Có điều kiện", "Email/Text", "Người quyết định ẩn.", "su_kien_thieu_hang.nguoi_an"),
        ),
    ),
    SheetSpec(
        "XAC_NHAN_THANG",
        "Xác nhận tháng không thiếu hàng",
        "Một dòng cho khoa × tháng khi không có sự kiện thiếu hàng; im lặng cũng phải được ghi nhận.",
        "xac_nhan_thang",
        "Cần để phân biệt không thiếu và không báo",
        (
            f("don_vi", "Bắt buộc", "Text", "Khoa xác nhận.", "xac_nhan_thang.don_vi"),
            f("thang", "Bắt buộc", "Integer", "Tháng xác nhận.", "xac_nhan_thang.thang", "1–12"),
            f("nam", "Bắt buộc", "Integer", "Năm xác nhận.", "xac_nhan_thang.nam"),
            f("khong_thieu", "Bắt buộc", "Boolean", "TRUE nghĩa là khoa xác nhận tháng đó không thiếu.", "xac_nhan_thang.khong_thieu", BOOL, "TRUE"),
            f("nguoi_xac_nhan", "Bắt buộc", "Email/Text", "Người chịu trách nhiệm xác nhận.", "xac_nhan_thang.nguoi_xac_nhan"),
            f("created_at", "Khuyến nghị", "Datetime", "Thời điểm xác nhận.", "xac_nhan_thang.created_at", DATETIME),
        ),
    ),
    SheetSpec(
        "SU_KIEN_NHU_CAU",
        "Sự kiện làm tăng/giảm nhu cầu",
        "Mỗi sự kiện phải được định lượng bằng %, số lượng/tháng hoặc số ca × định mức.",
        "su_kien_nhu_cau",
        "Cần để công thức phản ánh thay đổi thực tế",
        (
            f("ma_su_kien_nguon", "Khuyến nghị", "Text", "Mã đối soát nguồn.", "Khóa ánh xạ import"),
            f("don_vi", "Bắt buộc", "Text", "Khoa chịu ảnh hưởng.", "su_kien_nhu_cau.don_vi"),
            f("ma_quan_ly", "Có điều kiện", "Text", "Nhóm kỹ thuật chịu ảnh hưởng.", "su_kien_nhu_cau.ma_quan_ly"),
            f("ma_hang", "Có điều kiện", "Text", "Mã hàng chịu ảnh hưởng; phải có ma_hang hoặc ma_quan_ly.", "su_kien_nhu_cau.ma_hang"),
            f("ma_ly_do", "Bắt buộc", "Text", "Mã lý do chuẩn trong MA_LY_DO.", "su_kien_nhu_cau.ma_ly_do"),
            f("tu_thang", "Bắt buộc", "Integer", "Tháng bắt đầu ảnh hưởng.", "su_kien_nhu_cau.tu_thang", "1–12"),
            f("tu_nam", "Bắt buộc", "Integer", "Năm bắt đầu.", "su_kien_nhu_cau.tu_nam"),
            f("den_thang", "Khuyến nghị", "Integer", "Tháng kết thúc; trống nếu chưa xác định.", "su_kien_nhu_cau.den_thang", "1–12"),
            f("den_nam", "Khuyến nghị", "Integer", "Năm kết thúc.", "su_kien_nhu_cau.den_nam"),
            f("cach_dinh_luong", "Bắt buộc", "Enum", "Phương pháp định lượng.", "su_kien_nhu_cau.cach_dinh_luong", "phan_tram | sl_thang | ca_x_dinh_muc"),
            f("gia_tri", "Có điều kiện", "Number", "% hoặc SL/tháng; bắt buộc trừ ca_x_dinh_muc.", "su_kien_nhu_cau.gia_tri"),
            f("so_ca_thang", "Có điều kiện", "Number", "Bắt buộc khi ca_x_dinh_muc.", "su_kien_nhu_cau.so_ca_thang", "> 0"),
            f("dinh_muc_ca", "Có điều kiện", "Number", "Số vật tư/ca; bắt buộc khi ca_x_dinh_muc.", "su_kien_nhu_cau.dinh_muc_ca", "> 0"),
            f("lo_trinh", "Tùy chọn", "Text", "Lộ trình tăng/giảm theo thời gian.", "su_kien_nhu_cau.lo_trinh", example="25% -> 50% -> 100%"),
            f("muc_chac_chan", "Bắt buộc", "Enum", "Mức trưởng thành của kế hoạch.", "su_kien_nhu_cau.muc_chac_chan", "y_tuong | du_kien | da_phe_duyet | dang_chay"),
            f("ma_bi_thay_the", "Có điều kiện", "Text", "Mã bị thay thế khi lý do liên quan đổi mã.", "su_kien_nhu_cau.ma_bi_thay_the"),
            f("bang_chung", "Khuyến nghị", "Text/Path", "Số quyết định, kế hoạch hoặc tên file bằng chứng.", "su_kien_nhu_cau.bang_chung"),
            f("nguoi_khai", "Bắt buộc", "Email/Text", "Người khai sự kiện.", "su_kien_nhu_cau.nguoi_khai"),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái xét duyệt.", "su_kien_nhu_cau.trang_thai", "cho_duyet | da_duyet | tu_choi"),
            f("ly_do_tu_choi", "Có điều kiện", "Text", "Bắt buộc khi trạng thái tu_choi.", "su_kien_nhu_cau.ly_do_tu_choi"),
            f("nguoi_duyet", "Tùy chọn", "Email/Text", "Người duyệt.", "su_kien_nhu_cau.nguoi_duyet"),
            f("ngay_duyet", "Tùy chọn", "Datetime", "Thời gian duyệt.", "su_kien_nhu_cau.ngay_duyet", DATETIME),
            f("an_khoi_bao_cao", "Bắt buộc", "Boolean", "Ẩn nhưng không xóa sự kiện.", "su_kien_nhu_cau.an_khoi_bao_cao", BOOL, "FALSE"),
        ),
    ),
    SheetSpec(
        "MOC_CAM_KET",
        "Ngưỡng cam kết sử dụng",
        "Các mốc % tối thiểu dùng để cảnh báo tiến độ sử dụng.",
        "moc_cam_ket_su_dung",
        "Cấu hình",
        (
            f("thang_thu", "Bắt buộc", "Integer", "Số tháng kể từ hàng về đợt đầu.", "moc_cam_ket_su_dung.thang_thu", "> 0", "6"),
            f("ty_le_toi_thieu", "Bắt buộc", "Decimal", "Tỷ lệ 0–1, ví dụ 0.20 là 20%.", "moc_cam_ket_su_dung.ty_le_toi_thieu", "0–1", "0.20"),
            f("ghi_chu", "Khuyến nghị", "Text", "Nhãn giải thích mốc.", "moc_cam_ket_su_dung.ghi_chu"),
        ),
    ),
    SheetSpec(
        "HO_SO_CU_MANIFEST",
        "Manifest hồ sơ Word/Excel cũ",
        "Mỗi dòng mô tả một file gốc. Gửi kèm thư mục file, không nhúng file vào workbook.",
        "lan_xuat_ho_so; có thể tạo snapshot cho ho_so_cong_tac",
        "Bắt buộc để hiển thị lịch sử hồ sơ cũ",
        (
            f("duong_dan_tuong_doi", "Bắt buộc", "Text", "Đường dẫn tính từ thư mục ho_so_cu, dùng dấu /.", "Khóa nối file", example="2025/Khoa_GMHS/cam-ket.docx"),
            f("ten_file", "Bắt buộc", "Text", "Tên file khớp chính xác file gốc.", "Metadata import"),
            f("loai_file", "Bắt buộc", "Enum", "Loại tài liệu.", "lan_xuat_ho_so.ma_ho_so/renderer", "word | excel"),
            f("ma_ho_so", "Bắt buộc", "Enum", "Mã mẫu hồ sơ của hệ thống.", "lan_xuat_ho_so.ma_ho_so", "chi_dinh_thau | cam_ket_sl | danh_muc_dvsd | de_nghi_mua | tong_hop_thau"),
            f("ten_ho_so", "Bắt buộc", "Text", "Tên hiển thị của hồ sơ.", "lan_xuat_ho_so.ten_ho_so"),
            f("loai_mua_sam", "Bắt buộc", "Enum", "Gói mua sắm của hồ sơ.", "lan_xuat_ho_so.loai_mua_sam", LOAI_MUA),
            f("ma_dot_nguon", "Khuyến nghị", "Text", "Khớp DOT_DE_XUAT; để trống nếu không xác định được.", "lan_xuat_ho_so.dot_id qua ánh xạ"),
            f("ten_dot_cu", "Khuyến nghị", "Text", "Tên đợt ghi trên file cũ để đối chiếu.", "Metadata import"),
            f("don_vi", "Có điều kiện", "Text", "Khoa sở hữu; dùng Toàn viện cho hồ sơ tổng hợp.", "lan_xuat_ho_so.don_vi"),
            f("so_dong", "Khuyến nghị", "Integer", "Số dòng danh mục trong file.", "lan_xuat_ho_so.so_dong", ">= 0"),
            f("ngay_xuat", "Bắt buộc", "Datetime", "Ngày giờ xuất/ban hành file.", "lan_xuat_ho_so.ngay_xuat", DATETIME),
            f("nguoi_xuat", "Bắt buộc", "Email/Text", "Nhân viên xuất/ban hành.", "lan_xuat_ho_so.nguoi_xuat"),
            f("nguoi_duyet", "Tùy chọn", "Email/Text", "Người duyệt nếu có.", "Snapshot noi_dung"),
            f("revision", "Khuyến nghị", "Integer", "Số phiên bản nếu bệnh viện quản lý.", "ho_so_cong_tac.revision", ">= 1"),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái tại thời điểm lưu.", "ho_so_cong_tac.trang_thai/snapshot", "ban_nhap | cho_pdd | dang_xet_duyet | pdd_da_sua | tu_choi | da_duyet | da_ban_hanh"),
            f("ghi_chu", "Tùy chọn", "Text", "Ghi chú nguồn hoặc tình trạng file.", "Snapshot noi_dung"),
            f("sha256", "Khuyến nghị", "Text", "Mã băm SHA-256 để phát hiện file trùng/sai.", "Metadata import"),
        ),
    ),
    SheetSpec(
        "BIEU_MAU_CHINH_THUC",
        "Danh mục biểu mẫu chính thức",
        "Cung cấp file trống mới nhất và một file đã điền mẫu cho mỗi loại.",
        "bieu_mau + renderer frontend",
        "Bắt buộc để xuất đúng mẫu bệnh viện",
        (
            f("ma_bieu_mau", "Bắt buộc", "Text", "Mã ổn định không dấu.", "bieu_mau.ma"),
            f("ten", "Bắt buộc", "Text", "Tên hiển thị trong bước chọn biểu mẫu.", "bieu_mau.ten"),
            f("ma_ho_so", "Bắt buộc", "Enum", "Loại hồ sơ mà mẫu sinh ra.", "Renderer frontend", "chi_dinh_thau | cam_ket_sl | danh_muc_dvsd | de_nghi_mua | tong_hop_thau"),
            f("loai_file", "Bắt buộc", "Enum", "Định dạng mẫu.", "Renderer frontend", "word | excel"),
            f("version_mau", "Bắt buộc", "Text", "Phiên bản/quyết định ban hành mẫu.", "Metadata mẫu"),
            f("hieu_luc_tu", "Khuyến nghị", "Date", "Ngày bắt đầu dùng mẫu.", "Metadata mẫu", DATE),
            f("hieu_luc_den", "Tùy chọn", "Date", "Ngày hết hiệu lực; trống nếu hiện hành.", "Metadata mẫu", DATE),
            f("ten_file_trong", "Bắt buộc", "Text", "Tên file mẫu trống gửi kèm.", "Tệp nguồn"),
            f("ten_file_da_dien", "Khuyến nghị", "Text", "Tên file mẫu đã điền đã xóa dữ liệu nhạy cảm.", "Tệp đối chiếu"),
            f("dang_dung", "Bắt buộc", "Boolean", "TRUE nếu là mẫu được chọn hiện tại.", "Metadata mẫu", BOOL),
            f("mo_ta", "Tùy chọn", "Text", "Quy tắc merge ô, chữ ký, đánh số hoặc lưu ý.", "bieu_mau.mo_ta"),
        ),
    ),
    SheetSpec(
        "DE_NGHI_SUA_TIEU_CHI",
        "Lịch sử đề nghị sửa tiêu chí",
        "Dữ liệu lịch sử các yêu cầu sửa danh mục; không dùng sheet này để sửa thẳng danh mục.",
        "de_nghi_sua_tieu_chi",
        "Lịch sử tùy chọn",
        (
            f("ma_de_nghi_nguon", "Khuyến nghị", "Text", "Mã đối soát yêu cầu.", "Khóa ánh xạ import"),
            f("cap", "Bắt buộc", "Enum", "Sửa cấp mã hàng hay mã quản lý.", "de_nghi_sua_tieu_chi.cap", "ma_hang | ma_quan_ly"),
            f("ma_hang", "Có điều kiện", "Text", "Bắt buộc khi cap=ma_hang.", "de_nghi_sua_tieu_chi.ma_hang"),
            f("ma_quan_ly", "Bắt buộc", "Text", "Mã quản lý liên quan.", "de_nghi_sua_tieu_chi.ma_quan_ly"),
            f("don_vi", "Bắt buộc", "Text", "Khoa đề nghị.", "de_nghi_sua_tieu_chi.don_vi"),
            f("noi_dung_cu_json", "Bắt buộc", "JSON/Text", "Snapshot các trường trước khi sửa.", "de_nghi_sua_tieu_chi.noi_dung_cu"),
            f("noi_dung_moi_json", "Bắt buộc", "JSON/Text", "Snapshot các trường đề nghị thay đổi.", "de_nghi_sua_tieu_chi.noi_dung_moi"),
            f("ly_do", "Khuyến nghị", "Text", "Lý do đề nghị sửa.", "de_nghi_sua_tieu_chi.ly_do"),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái xử lý.", "de_nghi_sua_tieu_chi.trang_thai", "cho_duyet | da_duyet | tu_choi"),
            f("ly_do_tu_choi", "Có điều kiện", "Text", "Bắt buộc khi trạng thái tu_choi.", "de_nghi_sua_tieu_chi.ly_do_tu_choi"),
            f("nguoi_de_nghi", "Bắt buộc", "Email/Text", "Người đề nghị.", "de_nghi_sua_tieu_chi.nguoi_de_nghi"),
            f("ngay_de_nghi", "Bắt buộc", "Datetime", "Thời điểm đề nghị.", "de_nghi_sua_tieu_chi.ngay_de_nghi", DATETIME),
            f("nguoi_duyet", "Tùy chọn", "Email/Text", "Người xử lý.", "de_nghi_sua_tieu_chi.nguoi_duyet"),
            f("ngay_duyet", "Tùy chọn", "Datetime", "Thời điểm xử lý.", "de_nghi_sua_tieu_chi.ngay_duyet", DATETIME),
        ),
    ),
    SheetSpec(
        "TON_KHO_HANG_VE",
        "Tồn kho, đơn đặt hàng và lịch giao",
        "Dữ liệu bổ sung để tính nhu cầu mua ròng, cảnh báo sắp hết và lead time. Schema hiện tại chưa lưu đủ các cột này.",
        "Cần thiết kế bảng tồn kho/lô giao mới",
        "Mở rộng để công thức đầy đủ",
        (
            f("ngay_chot", "Bắt buộc", "Date", "Ngày chụp số liệu tồn kho.", "Bảng mới", DATE),
            f("kho", "Bắt buộc", "Text", "Kho sở hữu tồn.", "Bảng mới"),
            f("don_vi", "Tùy chọn", "Text", "Khoa sở hữu nếu là kho vệ tinh.", "Bảng mới"),
            f("ma_hang", "Bắt buộc", "Text", "Mã hàng.", "Bảng mới"),
            f("ton_dung_duoc", "Bắt buộc", "Number", "Tồn khả dụng có thể cấp.", "Bảng mới", ">= 0"),
            f("ton_bi_giu", "Bắt buộc", "Number", "Tồn đã giữ/đã phân bổ, chưa thể dùng tự do.", "Bảng mới", ">= 0"),
            f("ton_het_han_hong", "Bắt buộc", "Number", "Tồn hết hạn/hỏng/cách ly.", "Bảng mới", ">= 0"),
            f("han_dung_lo", "Khuyến nghị", "Date", "Hạn dùng của lô gần nhất; nếu nhiều lô ghi nhiều dòng.", "Bảng mới", DATE),
            f("so_hop_dong", "Khuyến nghị", "Text", "Số hợp đồng/đơn đặt hàng.", "Bảng mới"),
            f("nha_cung_cap", "Khuyến nghị", "Text", "Nhà cung cấp.", "Bảng mới"),
            f("don_gia", "Khuyến nghị", "Number", "Đơn giá chưa/đã thuế phải ghi rõ ở ghi_chu.", "Bảng mới", ">= 0"),
            f("tien_te", "Khuyến nghị", "Text", "Mã tiền tệ, mặc định VND.", "Bảng mới", "VND | USD | EUR | khác", "VND"),
            f("so_luong_da_dat", "Bắt buộc", "Number", "Số lượng đã đặt nhưng chưa chắc chắn về.", "Bảng mới", ">= 0"),
            f("so_luong_chac_chan_ve", "Bắt buộc", "Number", "Phần đã được nhà cung cấp xác nhận giao.", "Bảng mới", ">= 0"),
            f("ngay_dat", "Khuyến nghị", "Date", "Ngày đặt hàng để tính lead time.", "Bảng mới", DATE),
            f("ngay_du_kien_ve", "Khuyến nghị", "Date", "Ngày hẹn giao.", "Bảng mới", DATE),
            f("ngay_thuc_te_ve", "Tùy chọn", "Date", "Ngày nhận thực tế.", "Bảng mới", DATE),
            f("so_luong_thuc_nhan", "Tùy chọn", "Number", "Số lượng thực nhận ở lần giao.", "Bảng mới", ">= 0"),
            f("trang_thai_giao", "Khuyến nghị", "Enum", "Trạng thái lô giao.", "Bảng mới", "chua_xac_nhan | da_xac_nhan | giao_mot_phan | da_giao | huy"),
            f("ghi_chu", "Tùy chọn", "Text", "Điều kiện giá, chậm giao hoặc ngoại lệ.", "Bảng mới"),
        ),
    ),
    SheetSpec(
        "QUY_DOI_THAY_THE",
        "Quy đổi đơn vị, đổi mã và mã thay thế",
        "Giữ lịch sử thay đổi mã/quy cách để số liệu các năm so sánh được.",
        "Cần bảng chuẩn hóa mới",
        "Mở rộng chất lượng dữ liệu",
        (
            f("loai_quan_he", "Bắt buộc", "Enum", "Loại bản ghi chuẩn hóa.", "Bảng mới", "quy_doi_dvt | doi_ma | thay_the_lam_sang"),
            f("ma_hang_nguon", "Bắt buộc", "Text", "Mã gốc/cũ.", "Bảng mới"),
            f("ma_hang_dich", "Có điều kiện", "Text", "Mã mới hoặc mã thay thế.", "Bảng mới"),
            f("dvt_nguon", "Có điều kiện", "Text", "ĐVT của mã nguồn.", "Bảng mới"),
            f("dvt_chuan", "Có điều kiện", "Text", "ĐVT quy đổi đích.", "Bảng mới"),
            f("he_so_quy_doi", "Có điều kiện", "Number", "1 ĐVT nguồn bằng bao nhiêu ĐVT chuẩn.", "Bảng mới", "> 0"),
            f("quy_cach_dong_goi", "Khuyến nghị", "Text", "Mô tả hộp/gói/thùng và số đơn vị.", "Bảng mới"),
            f("hieu_luc_tu", "Bắt buộc", "Date", "Ngày bắt đầu áp dụng quan hệ.", "Bảng mới", DATE),
            f("hieu_luc_den", "Tùy chọn", "Date", "Ngày ngừng áp dụng.", "Bảng mới", DATE),
            f("ma_nhom_thay_the", "Có điều kiện", "Text", "Nhóm các mã tương đương lâm sàng.", "Bảng mới"),
            f("uu_tien", "Tùy chọn", "Integer", "Thứ tự ưu tiên dùng mã thay thế.", "Bảng mới", ">= 1"),
            f("bang_chung", "Khuyến nghị", "Text/Path", "Hội đồng/phê duyệt chuyên môn cho phép thay thế.", "Bảng mới"),
            f("ghi_chu", "Tùy chọn", "Text", "Điều kiện chỉ định hoặc ngoại lệ.", "Bảng mới"),
        ),
    ),
    SheetSpec(
        "MA_LY_DO",
        "Danh mục mã lý do",
        "Danh mục đóng dùng chung cho thiếu hàng và sự kiện nhu cầu.",
        "ma_ly_do",
        "Cấu hình",
        (
            f("ma", "Bắt buộc", "Text", "Mã lý do duy nhất.", "ma_ly_do.ma", example="A1"),
            f("nhom", "Bắt buộc", "Enum", "A thiếu/cấp hạn chế; B sai lệch ghi nhận; C thay đổi nhu cầu; D chất lượng dữ liệu.", "ma_ly_do.nhom", "A | B | C | D"),
            f("ten", "Bắt buộc", "Text", "Tên ngắn hiển thị.", "ma_ly_do.ten"),
            f("mo_ta", "Khuyến nghị", "Text", "Định nghĩa và ví dụ sử dụng.", "ma_ly_do.mo_ta"),
            f("dang_dung", "Bắt buộc", "Boolean", "Không xóa mã đã dùng; chuyển FALSE khi ngừng.", "ma_ly_do.dang_dung", BOOL),
            f("thu_tu", "Bắt buộc", "Integer", "Thứ tự hiển thị trong nhóm.", "ma_ly_do.thu_tu", ">= 0"),
        ),
    ),
    SheetSpec(
        "MA_MOI_KHOA_DE_NGHI",
        "Lịch sử mã mới khoa đề nghị",
        "Dữ liệu các yêu cầu bổ sung mã/nhóm mới của khoa.",
        "khoa_nhom_ky_thuat",
        "Lịch sử tùy chọn",
        (
            f("ma_de_nghi_nguon", "Khuyến nghị", "Text", "Mã đối soát.", "Khóa ánh xạ import"),
            f("don_vi", "Bắt buộc", "Text", "Khoa đề nghị.", "khoa_nhom_ky_thuat.don_vi"),
            f("ma_quan_ly", "Tùy chọn", "Text", "Mã quản lý nếu đã biết.", "khoa_nhom_ky_thuat.ma_quan_ly"),
            f("la_nhom_moi", "Bắt buộc", "Boolean", "TRUE nếu tạo nhóm/mã mới.", "khoa_nhom_ky_thuat.la_nhom_moi", BOOL),
            f("ten_quan_ly_moi", "Có điều kiện", "Text", "Tên nhóm mới.", "khoa_nhom_ky_thuat.ten_quan_ly_moi"),
            f("ma_hang_moi", "Tùy chọn", "Text", "Mã mới; hệ thống có thể sinh MOI-<id> nếu trống.", "khoa_nhom_ky_thuat.ma_hang_moi"),
            f("ten_vat_tu_moi", "Bắt buộc", "Text", "Tên vật tư mới.", "khoa_nhom_ky_thuat.ten_vat_tu_moi"),
            f("dvt_moi", "Bắt buộc", "Text", "Đơn vị tính.", "khoa_nhom_ky_thuat.dvt_moi"),
            f("ten_thuong_mai", "Tùy chọn", "Text", "Tên thương mại.", "khoa_nhom_ky_thuat.ten_thuong_mai"),
            f("tieu_chi_ky_thuat", "Khuyến nghị", "Text", "Tiêu chí kỹ thuật.", "khoa_nhom_ky_thuat.tieu_chi_ky_thuat"),
            f("ky_ma_hieu", "Tùy chọn", "Text", "Model/catalogue.", "khoa_nhom_ky_thuat.ky_ma_hieu"),
            f("hang", "Tùy chọn", "Text", "Hãng.", "khoa_nhom_ky_thuat.hang"),
            f("nuoc_san_xuat", "Tùy chọn", "Text", "Nước sản xuất.", "khoa_nhom_ky_thuat.nuoc_san_xuat"),
            f("goi", "Khuyến nghị", "Text", "Gói chuyên môn.", "khoa_nhom_ky_thuat.goi"),
            f("so_luong", "Bắt buộc", "Number", "Số lượng đề xuất.", "khoa_nhom_ky_thuat.so_luong", ">= 0"),
            f("tu_thang", "Bắt buộc", "Integer", "Tháng bắt đầu.", "khoa_nhom_ky_thuat.tu_thang", "1–12"),
            f("tu_nam", "Bắt buộc", "Integer", "Năm bắt đầu.", "khoa_nhom_ky_thuat.tu_nam"),
            f("den_thang", "Bắt buộc", "Integer", "Tháng kết thúc.", "khoa_nhom_ky_thuat.den_thang", "1–12"),
            f("den_nam", "Bắt buộc", "Integer", "Năm kết thúc.", "khoa_nhom_ky_thuat.den_nam"),
            f("ghi_chu", "Tùy chọn", "Text", "Giải trình.", "khoa_nhom_ky_thuat.ghi_chu"),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái duyệt.", "khoa_nhom_ky_thuat.trang_thai", "cho_duyet | da_duyet | tu_choi"),
            f("ly_do_tu_choi", "Có điều kiện", "Text", "Bắt buộc nếu từ chối.", "khoa_nhom_ky_thuat.ly_do_tu_choi"),
            f("created_by", "Bắt buộc", "Email", "Người đề nghị.", "khoa_nhom_ky_thuat.created_by"),
            f("created_by_ho_ten", "Khuyến nghị", "Text", "Họ tên snapshot.", "khoa_nhom_ky_thuat.created_by_ho_ten"),
            f("created_at", "Bắt buộc", "Datetime", "Thời điểm đề nghị.", "khoa_nhom_ky_thuat.created_at", DATETIME),
            f("duyet_boi", "Tùy chọn", "Email/Text", "Người duyệt.", "khoa_nhom_ky_thuat.duyet_boi"),
            f("duyet_luc", "Tùy chọn", "Datetime", "Thời gian duyệt.", "khoa_nhom_ky_thuat.duyet_luc", DATETIME),
        ),
    ),
    SheetSpec(
        "GAN_NHOM_VAO_GOI",
        "Gán nhóm kỹ thuật vào gói thầu",
        "Một mã quản lý chỉ thuộc một gói trong cùng năm.",
        "goi_thau + goi_thau_assignment",
        "Cần cho tổng hợp theo gói",
        (
            f("ma_goi_danh_muc", "Bắt buộc", "Text", "Mã đối soát gói danh mục.", "Khóa ánh xạ import"),
            f("ten_goi", "Bắt buộc", "Text", "Tên gói.", "goi_thau.ten_goi"),
            f("nam", "Bắt buộc", "Integer", "Năm gói và năm gán.", "goi_thau.nam / goi_thau_assignment.nam"),
            f("trang_thai", "Bắt buộc", "Enum", "Trạng thái gói danh mục.", "goi_thau.trang_thai", "dang_lap | da_trinh | da_trung"),
            f("loai_mua_sam", "Bắt buộc", "Enum", "Loại mua sắm.", "goi_thau.loai_mua_sam", f"{LOAI_MUA} | khac"),
            f("ma_quan_ly", "Bắt buộc", "Text", "Nhóm được gán.", "goi_thau_assignment.ma_quan_ly"),
            f("assigned_by", "Khuyến nghị", "Email/Text", "Người gán.", "goi_thau_assignment.assigned_by"),
            f("assigned_at", "Khuyến nghị", "Datetime", "Thời điểm gán.", "goi_thau_assignment.assigned_at", DATETIME),
        ),
    ),
)


STAGING_COUNTS = (
    ("nhom_ky_thuat", 878, "Đã có"),
    ("vat_tu", 3085, "Đã có"),
    ("users", 7, "Cần danh sách đầy đủ"),
    ("usage_history_current", 149999, "Đã có 1 mẻ; cần cập nhật định kỳ"),
    ("proposals", 16, "Chỉ có ít dữ liệu thử/lịch sử"),
    ("proposal_reasons", 16, "Đi kèm đề xuất hiện có"),
    ("dot_de_xuat", 5, "Đã có"),
    ("goi_thau_tien_do", 1, "Thiếu lịch sử các gói"),
    ("goi_thau_moc", 5, "Chỉ đủ mốc cho 1 gói"),
    ("goi_thau_ket_qua_ma", 13, "Thiếu kết quả lịch sử toàn viện"),
    ("ky_thau", 0, "Thiếu hoàn toàn"),
    ("so_luong_ky", 0, "Thiếu hoàn toàn"),
    ("hop_dong", 0, "Thiếu hoàn toàn"),
    ("su_kien_thieu_hang", 2, "Mới có dữ liệu thử"),
    ("xac_nhan_thang", 1, "Mới có dữ liệu thử"),
    ("su_kien_nhu_cau", 1, "Mới có dữ liệu thử"),
    ("moc_cam_ket_su_dung", 3, "Đã có cấu hình 6/12/18 tháng"),
    ("ho_so_cong_tac", 4, "Mới có vài hồ sơ"),
    ("lan_xuat_ho_so", 13, "Chưa phản ánh kho hồ sơ cũ của bệnh viện"),
)


BASE_FILES = (
    "backend/sql/schema.sql",
    "backend/sql/rls_policies.sql",
)

PATCH_FILES = (
    "patch_a2_cong_phe_duyet.sql",
    "patch_a4_tien_do_goi_thau.sql",
    "patch_a5_danh_sach_khoa.sql",
    "patch_bc_so_goc_va_so_ghi.sql",
    "patch_h_dot_va_lich_su_xuat.sql",
    "patch_i_rut_va_tong_hop.sql",
    "patch_j_dong_bo_sequence_staging.sql",
    "patch_k_ho_so_cong_tac_truc_tuyen.sql",
    "patch_l_gio_nhap_tren_server.sql",
    "patch_m_ket_qua_thau_ve_khoa.sql",
    "patch_n_dieu_chinh_tieu_chi.sql",
    "patch_o_tien_do_su_dung.sql",
    "patch_p_du_kien_het_hang.sql",
    "patch_q_phan_nhom_abc.sql",
    "patch_r_nhu_cau_bi_nen.sql",
    "patch_s_workflow_ho_so_dvsd.sql",
    "patch_t_tao_nhieu_bo_ho_so.sql",
)


NAVY = "0B2B4C"
BLUE = "DCEEFF"
TEAL = "DDF7F2"
ORANGE = "FDE9D2"
YELLOW = "FFF4CC"
GRAY = "E8EDF3"
LIGHT = "F7FAFC"
RED = "FCE1E1"
WHITE = "FFFFFF"
GREEN = "DFF2E1"
TEXT = "1E293B"
THIN = Side(style="thin", color="CBD5E1")


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def style_title(cell, size=18):
    cell.font = Font(bold=True, color=WHITE, size=size)
    cell.fill = fill(NAVY)
    cell.alignment = Alignment(vertical="center")


def style_header(cell, level: str = ""):
    color = {
        "Bắt buộc": ORANGE,
        "Có điều kiện": YELLOW,
        "Khuyến nghị": BLUE,
        "Tùy chọn": GRAY,
    }.get(level, NAVY)
    cell.fill = fill(color)
    cell.font = Font(bold=True, color=TEXT if color != NAVY else WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add_note(cell, field: Field):
    parts = [
        f"Mức độ: {field.level}",
        f"Kiểu dữ liệu: {field.dtype}",
        f"Đích: {field.target}",
        f"Ý nghĩa: {field.description}",
    ]
    if field.allowed:
        parts.append(f"Giá trị hợp lệ: {field.allowed}")
    if field.example:
        parts.append(f"Ví dụ: {field.example}")
    cell.comment = Comment("\n".join(parts), "Codex")


def add_validation(ws, col_idx: int, field: Field):
    start, end = 2, 10001
    letter = get_column_letter(col_idx)
    allowed = [x.strip() for x in field.allowed.split("|") if x.strip()]
    allowed = [x for x in allowed if x not in {"trống", "khác"}]
    if allowed and len(",".join(allowed)) < 240 and all("–" not in x and ".." not in x and ">=" not in x and ">" not in x for x in allowed):
        dv = DataValidation(type="list", formula1=f'"{",".join(allowed)}"', allow_blank=field.level != "Bắt buộc")
        dv.error = "Giá trị không thuộc danh mục cho phép."
        dv.errorTitle = "Giá trị không hợp lệ"
        dv.prompt = field.allowed
        dv.promptTitle = field.name
        dv.showErrorMessage = True
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{letter}{start}:{letter}{end}")
    elif field.dtype in {"Integer", "Number", "Decimal"} and (">= 0" in field.allowed or "> 0" in field.allowed):
        op = "greaterThan" if "> 0" in field.allowed and ">= 0" not in field.allowed else "greaterThanOrEqual"
        dv = DataValidation(type="decimal", operator=op, formula1="0", allow_blank=field.level != "Bắt buộc")
        dv.error = "Phải nhập số không âm/dương theo mô tả cột."
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{letter}{start}:{letter}{end}")
    elif field.name in {"thang", "tu_thang", "den_thang", "thang_moc"}:
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=field.level != "Bắt buộc")
        ws.add_data_validation(dv)
        dv.add(f"{letter}{start}:{letter}{end}")
    elif field.dtype == "Date":
        dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)", allow_blank=field.level != "Bắt buộc")
        ws.add_data_validation(dv)
        dv.add(f"{letter}{start}:{letter}{end}")


def add_guide(wb: Workbook):
    ws = wb.active
    ws.title = "00_HUONG_DAN"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "MẪU DỮ LIỆU BỆNH VIỆN — FULL WORKFLOW VTYT"
    style_title(ws["A1"], 20)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:H2")
    ws["A2"] = "Phiên bản lập ngày 03/08/2026 · Dùng để kiểm tra, làm sạch và nạp STAGING trước khi đưa lên production"
    ws["A2"].fill = fill(BLUE)
    ws["A2"].font = Font(italic=True, color=TEXT)

    notes = [
        ("1. Database staging hiện lấy từ đâu?", "Không lấy từ một file Excel duy nhất. Cấu trúc được dựng bằng schema.sql + rls_policies.sql, sau đó áp các patch A2/A4/A5/BC/H→T. Dữ liệu lõi ban đầu được sao chép production → JSON → staging bằng backend/scripts/xuat_du_lieu_sang_staging.py."),
        ("2. Cách điền workbook", "Mỗi sheet dữ liệu có header ở dòng 1; bắt đầu nhập từ dòng 2. Không đổi tên sheet/cột. Giữ mọi mã ở dạng text, ngày theo YYYY-MM-DD, số không kèm đơn vị. Xem toàn bộ giải thích tại sheet 01_TU_DIEN_COT hoặc rê chuột vào header."),
        ("3. Màu cột", "Cam = bắt buộc; vàng = có điều kiện; xanh = khuyến nghị; xám = tùy chọn. Một số cột mở rộng chưa có bảng DB tương ứng; vẫn nên cung cấp để thiết kế schema mà không phải xin lại dữ liệu."),
        ("4. Tệp hồ sơ", "Không nhúng Word/Excel cũ vào workbook. Gửi riêng thư mục ho_so_cu/ và khai báo từng file trong HO_SO_CU_MANIFEST. Không gửi mật khẩu, service key, token hoặc dữ liệu người bệnh."),
        ("5. Quy trình nạp", "Nhận file → báo cáo chất lượng → chuẩn hóa mã/tên khoa → nạp staging → test bằng tài khoản 3 role → đối soát tổng dòng/tổng lượng → chỉ khi nghiệm thu mới nạp production."),
    ]
    row = 4
    for title, detail in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1, title).font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row, 3, detail).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 48
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "NGUỒN CẤU TRÚC DATABASE STAGING").font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = fill(NAVY)
    row += 1
    for path in BASE_FILES:
        ws.cell(row, 1, "Baseline")
        ws.cell(row, 2, path)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1
    for path in PATCH_FILES:
        ws.cell(row, 1, "Patch")
        ws.cell(row, 2, f"backend/sql/{path}")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "SỐ DÒNG ĐỌC TRỰC TIẾP TỪ STAGING (03/08/2026)").font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = fill(NAVY)
    row += 1
    for c, val in enumerate(("Bảng", "Số dòng", "Nhận định"), 1):
        ws.cell(row, c, val)
        style_header(ws.cell(row, c))
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
    for table, count, status in STAGING_COUNTS:
        row += 1
        ws.cell(row, 1, table)
        ws.cell(row, 2, count)
        ws.cell(row, 3, status)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        if count == 0:
            ws.cell(row, 1).fill = ws.cell(row, 2).fill = ws.cell(row, 3).fill = fill(RED)

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "THỨ TỰ BÀN GIAO KHUYẾN NGHỊ").font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = fill(NAVY)
    for i, text in enumerate(
        (
            "1) DM_VAT_TU + LICH_SU_XUAT_KHO + DON_VI_TAI_KHOAN.",
            "2) DOT_DE_XUAT + DE_XUAT_CU + LY_DO_DE_XUAT + hồ sơ cũ.",
            "3) KY_THAU_SO_CHOT + GOI_THAU_TIMELINE + KET_QUA_THAU + HOP_DONG.",
            "4) THIEU_HANG + XAC_NHAN_THANG + SU_KIEN_NHU_CAU.",
            "5) TON_KHO_HANG_VE + QUY_DOI_THAY_THE và biểu mẫu chính thức.",
        ),
        1,
    ):
        ws.cell(row + i, 1, text)
        ws.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=8)

    widths = (24, 28, 22, 18, 18, 18, 18, 18)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A3"


def add_dictionary(wb: Workbook):
    ws = wb.create_sheet("01_TU_DIEN_COT")
    ws.sheet_view.showGridLines = False
    headers = ("Sheet", "Mục đích", "Cột", "Mức độ", "Kiểu dữ liệu", "Giá trị hợp lệ", "Đích database", "Giải thích chi tiết", "Ví dụ")
    ws.append(headers)
    for cell in ws[1]:
        style_header(cell)
    for spec in SHEETS:
        for field in spec.fields:
            ws.append((spec.name, spec.purpose, field.name, field.level, field.dtype, field.allowed, field.target, field.description, field.example))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E2E8F0"))
        row[3].fill = fill({
            "Bắt buộc": ORANGE,
            "Có điều kiện": YELLOW,
            "Khuyến nghị": BLUE,
            "Tùy chọn": GRAY,
        }.get(row[3].value, WHITE))
    widths = (24, 42, 25, 16, 16, 34, 34, 58, 25)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"


def add_data_sheet(wb: Workbook, spec: SheetSpec):
    ws = wb.create_sheet(spec.name)
    ws.sheet_view.showGridLines = False
    for idx, field in enumerate(spec.fields, 1):
        cell = ws.cell(1, idx, field.name)
        style_header(cell, field.level)
        add_note(cell, field)
        add_validation(ws, idx, field)
        width = max(13, min(34, max(len(field.name) + 3, len(field.example) + 3)))
        if field.dtype in {"Text dài", "JSON/Text", "Text/Path"}:
            width = 34
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.fields))}1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.center.text = f"&B{spec.title}"
    ws.oddFooter.left.text = f"Đích: {spec.target}"
    ws.oddFooter.right.text = "Trang &P/&N"

    # Tô đỏ dòng có dữ liệu nhưng thiếu trường bắt buộc đầu tiên, giúp người
    # dùng thấy lỗi sơ bộ ngay trong Excel mà không cần macro.
    required = [i for i, field in enumerate(spec.fields, 1) if field.level == "Bắt buộc"]
    if required:
        first = get_column_letter(required[0])
        last = get_column_letter(len(spec.fields))
        formula = f'=AND(COUNTA($A2:${last}2)>0,${first}2="")'
        ws.conditional_formatting.add(
            f"A2:{last}10001",
            FormulaRule(formula=[formula], fill=fill(RED)),
        )


def build(output: Path):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    add_guide(wb)
    add_dictionary(wb)
    for spec in SHEETS:
        add_data_sheet(wb, spec)
    wb.save(output)

    # Mở lại để xác nhận file ZIP/XML hợp lệ và đủ sheet/header.
    check = load_workbook(output, read_only=True, data_only=False)
    expected = {"00_HUONG_DAN", "01_TU_DIEN_COT", *(spec.name for spec in SHEETS)}
    missing = expected.difference(check.sheetnames)
    if missing:
        raise RuntimeError(f"Workbook thiếu sheet: {sorted(missing)}")
    for spec in SHEETS:
        actual = tuple(cell.value for cell in next(check[spec.name].iter_rows(min_row=1, max_row=1)))
        wanted = tuple(field.name for field in spec.fields)
        if actual != wanted:
            raise RuntimeError(f"Header sheet {spec.name} không khớp.")
    check.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "output",
        nargs="?",
        default="MAU_DU_LIEU_BENH_VIEN_FULL_WORKFLOW_VTYT.xlsx",
        type=Path,
    )
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
