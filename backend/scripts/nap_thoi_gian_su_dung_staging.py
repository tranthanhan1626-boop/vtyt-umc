#!/usr/bin/env python3
"""Nạp phần dữ liệu dùng được từ file thời gian sử dụng vào STAGING.

Mặc định chỉ dry-run. Thêm --apply mới ghi danh mục:
  - thêm nhóm kỹ thuật chưa có;
  - thêm mã hàng chưa có;
  - chỉ điền ma_quan_ly / ten_thuong_mai đang NULL ở mã hiện có.

Thêm --apply-snapshot để nạp ảnh chụp khả dụng/hợp đồng vào schema patch U.
Không ghi đè giá trị hiện hữu và không ghi production.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BACKUP_ROOT = ROOT / "backend" / "du_lieu_staging"
PAGE = 1000

COLS = {
    "Mã nhóm quản lý": "ma_quan_ly",
    "Tên nhóm quản lý": "ten_quan_ly",
    "Mã hàng": "ma_hang",
    "Tên hàng": "ten_vat_tu",
    "Tên thương mại": "ten_thuong_mai",
    "List TG theo Mã hàng - Số QĐ": "so_quyet_dinh",
    "ĐVT tồn kho": "dvt",
    "SL hợp đồng": "sl_hop_dong",
    "SL hợp đồng CS1": "sl_hop_dong_cs1",
    "SL chưa thực hiện hợp đồng CS1": "sl_chua_thuc_hien_hop_dong_cs1",
    "SL mua thêm 30%": "sl_mua_them_30",
    "SL đã mua thêm 30%": "sl_da_mua_them_30",
    "SL còn có thể mua thêm 30%": "sl_con_co_the_mua_them_30",
    "SL đã mua thêm chưa lãnh hàng": "sl_da_mua_them_chua_lanh",
    "SL đã thông qua hội đồng": "sl_da_thong_qua_hoi_dong",
    "SL đang chào giá": "sl_dang_chao_gia",
    "SL tồn CS1": "sl_ton_cs1",
    "SL khả dụng CS1": "sl_kha_dung_cs1",
    "SL khả dụng CS1 (30%)": "sl_kha_dung_cs1_30",
    "SL sd 2023": "sl_sd_2023",
    "SL sd 2024": "sl_sd_2024",
    "SL sd 2025": "sl_sd_2025",
    "SL sd 2026": "sl_sd_2026",
    "SL sử dụng trung bình": "sl_su_dung_trung_binh",
    "Thời gian đáp ứng (mã hàng)": "thoi_gian_dap_ung_ma_hang",
    "Thời gian đáp ứng (mql)": "thoi_gian_dap_ung_ma_quan_ly",
    "Thời gian đáp ứng (mã hàng) (30%)": "thoi_gian_dap_ung_ma_hang_30",
    "Thời gian đáp ứng (mql) (30%)": "thoi_gian_dap_ung_ma_quan_ly_30",
    "List Công ty theo Mã hàng": "nha_cung_cap",
}


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def text(value: Any) -> str | None:
    value = clean(value)
    return None if value is None else str(value)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def parse(path: Path):
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    headers = [text(v) or "" for v in all_rows[0]]
    unknown = [h for h in headers if h not in COLS]
    missing = [h for h in COLS if h not in headers]
    if unknown or missing:
        raise RuntimeError(f"Header khác dự kiến. Thừa={unknown}; thiếu={missing}")

    parsed, excluded = [], []
    for source_row, values in enumerate(all_rows[1:], 2):
        original = dict(zip(headers, values))
        code = text(original.get("Mã hàng"))
        first = text(values[0] if values else None)
        if not any(v not in (None, "") for v in values):
            excluded.append({"source_row": source_row, "reason": "blank"})
            continue
        if code == "Total" or first == "Total":
            excluded.append({"source_row": source_row, "reason": "total"})
            continue
        if first and first.startswith("Bộ lọc được áp dụng"):
            excluded.append({"source_row": source_row, "reason": "filter_note"})
            continue
        if not code:
            excluded.append({"source_row": source_row, "reason": "missing_ma_hang"})
            continue
        row = {"source_row": source_row}
        for header, key in COLS.items():
            row[key] = clean(original.get(header))
        row["ma_hang"] = text(row["ma_hang"])
        row["ma_quan_ly"] = text(row["ma_quan_ly"])
        row["ten_quan_ly"] = text(row["ten_quan_ly"])
        row["ten_vat_tu"] = text(row["ten_vat_tu"])
        row["ten_thuong_mai"] = text(row["ten_thuong_mai"])
        row["dvt"] = text(row["dvt"])
        parsed.append(row)
    wb.close()
    return parsed, excluded


class Rest:
    def __init__(self, url: str, key: str):
        self.base = url.rstrip("/") + "/rest/v1/"
        self.key = key

    def request(
        self,
        method: str,
        path: str,
        body: Any = None,
        prefer: str | None = None,
        extra_headers: dict[str, str] | None = None,
    ):
        data = None if body is None else json.dumps(body, ensure_ascii=False, default=str).encode()
        headers = {
            "apikey": self.key,
            "Authorization": "Bearer " + self.key,
            "Content-Type": "application/json",
        }
        if prefer:
            headers["Prefer"] = prefer
        if extra_headers:
            headers.update(extra_headers)
        req = urllib.request.Request(self.base + path, data=data, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=60) as response:
                raw = response.read()
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", "replace")
            raise RuntimeError(f"{method} {path}: HTTP {exc.code}: {detail[:800]}") from exc

    def fetch_all(self, table: str, select: str):
        out, start = [], 0
        while True:
            query = urllib.parse.urlencode({"select": select}, safe=",*")
            batch = self.request(
                "GET",
                f"{table}?{query}",
                extra_headers={"Range": f"{start}-{start + PAGE - 1}"},
            )
            out.extend(batch)
            if len(batch) < PAGE:
                return out
            start += PAGE

    def insert_ignore(self, table: str, conflict: str, rows: list[dict]):
        for start in range(0, len(rows), 200):
            chunk = rows[start : start + 200]
            query = urllib.parse.urlencode({"on_conflict": conflict})
            self.request(
                "POST",
                f"{table}?{query}",
                chunk,
                prefer="resolution=ignore-duplicates,return=minimal",
            )

    def insert_return(self, table: str, rows: list[dict], conflict: str | None = None):
        query = ""
        if conflict:
            query = "?" + urllib.parse.urlencode({"on_conflict": conflict})
        return self.request(
            "POST",
            f"{table}{query}",
            rows,
            prefer="resolution=ignore-duplicates,return=representation",
        )

    def patch_one(self, table: str, key_col: str, key_value: str, changes: dict):
        encoded = urllib.parse.quote(key_value, safe="")
        self.request(
            "PATCH",
            f"{table}?{key_col}=eq.{encoded}",
            changes,
            prefer="return=minimal",
        )


def singleton(rows: list[dict], key: str) -> Any:
    values = {row.get(key) for row in rows if row.get(key) not in (None, "")}
    return next(iter(values)) if len(values) == 1 else None


def analyze(source_rows: list[dict], groups: list[dict], items: list[dict]):
    by_code: dict[str, list[dict]] = defaultdict(list)
    for row in source_rows:
        by_code[row["ma_hang"]].append(row)

    current_groups = {row["ma_quan_ly"]: row for row in groups}
    current_items = {row["ma_hang"]: row for row in items}

    group_values: dict[str, set[str]] = defaultdict(set)
    for row in source_rows:
        if row.get("ma_quan_ly") and row.get("ten_quan_ly"):
            group_values[row["ma_quan_ly"]].add(row["ten_quan_ly"])
    ambiguous_groups = {code: sorted(names) for code, names in group_values.items() if len(names) > 1}
    new_groups = [
        {"ma_quan_ly": code, "ten_quan_ly": next(iter(names))}
        for code, names in group_values.items()
        if len(names) == 1 and code not in current_groups
    ]

    new_items, patches, conflicts, source_ambiguities = [], [], [], []
    now = datetime.now(timezone.utc).isoformat()
    for code, rows in by_code.items():
        values = {
            key: singleton(rows, key)
            for key in ("ma_quan_ly", "ten_vat_tu", "ten_thuong_mai", "dvt")
        }
        for key in ("ma_quan_ly", "ten_vat_tu", "ten_thuong_mai", "dvt"):
            distinct = sorted({str(row[key]) for row in rows if row.get(key)})
            if len(distinct) > 1:
                source_ambiguities.append({"ma_hang": code, "field": key, "values": distinct})

        current = current_items.get(code)
        if current is None:
            if not values["ten_vat_tu"]:
                conflicts.append({"ma_hang": code, "field": "ten_vat_tu", "reason": "missing_for_insert"})
                continue
            new_items.append(
                {
                    "ma_hang": code,
                    "ten_vat_tu": values["ten_vat_tu"],
                    "dvt": values["dvt"],
                    "ma_quan_ly": values["ma_quan_ly"],
                    "ten_thuong_mai": values["ten_thuong_mai"],
                }
            )
            continue

        changes = {}
        for field in ("ma_quan_ly", "ten_thuong_mai"):
            incoming = values[field]
            existing = current.get(field)
            if incoming and not existing:
                changes[field] = incoming
            elif incoming and existing and str(existing).strip() != str(incoming).strip():
                conflicts.append(
                    {
                        "ma_hang": code,
                        "field": field,
                        "staging": existing,
                        "source": incoming,
                        "reason": "do_not_overwrite",
                    }
                )
        if changes:
            changes["updated_at"] = now
            patches.append({"ma_hang": code, "changes": changes})

        for field in ("ten_vat_tu", "dvt"):
            incoming = values[field]
            existing = current.get(field)
            if incoming and existing and str(existing).strip() != str(incoming).strip():
                conflicts.append(
                    {
                        "ma_hang": code,
                        "field": field,
                        "staging": existing,
                        "source": incoming,
                        "reason": "informational_only",
                    }
                )

    negative_fields = Counter()
    for row in source_rows:
        for key, value in row.items():
            if key.startswith("sl_") and isinstance(value, (int, float)) and value < 0:
                negative_fields[key] += 1

    return {
        "by_code": by_code,
        "new_groups": new_groups,
        "new_items": new_items,
        "patches": patches,
        "conflicts": conflicts,
        "ambiguous_groups": ambiguous_groups,
        "source_ambiguities": source_ambiguities,
        "negative_fields": dict(negative_fields),
        "counts": {
            "source_rows": len(source_rows),
            "source_unique_items": len(by_code),
            "source_duplicate_items": sum(len(rows) > 1 for rows in by_code.values()),
            "staging_groups_before": len(groups),
            "staging_items_before": len(items),
            "new_groups": len(new_groups),
            "new_items": len(new_items),
            "existing_items_to_patch": len(patches),
            "fill_ma_quan_ly": sum("ma_quan_ly" in p["changes"] for p in patches),
            "fill_ten_thuong_mai": sum("ten_thuong_mai" in p["changes"] for p in patches),
            "conflicts_not_overwritten": len(conflicts),
        },
    }


def write_json(path: Path, value: Any):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def apply_snapshot(
    api: Rest,
    source: Path,
    digest: str,
    source_rows: list[dict],
    as_of_date: str | None,
):
    """Nạp nguyên từng dòng nguồn; không cộng các mã hàng trùng nhau."""
    header = {
        "ten_file": source.name,
        "sha256_file": digest,
        "ngay_chot_so": as_of_date,
        "ghi_chu": (
            "Nguồn thời gian sử dụng chi tiết có mua thêm 30%. "
            "Chưa có ngày hết hiệu lực hợp đồng trong file nguồn."
        ),
        "so_dong_nguon": len(source_rows),
        "so_ma_hang": len({r["ma_hang"] for r in source_rows}),
    }
    inserted = api.insert_return(
        "nguon_kha_dung_hop_dong", [header], conflict="sha256_file"
    )
    if inserted:
        source_id = inserted[0]["id"]
    else:
        encoded = urllib.parse.quote(digest, safe="")
        found = api.request(
            "GET",
            f"nguon_kha_dung_hop_dong?sha256_file=eq.{encoded}&select=id",
        )
        if not found:
            raise RuntimeError("Không lấy được id nguồn ảnh chụp sau insert.")
        source_id = found[0]["id"]

    payload = []
    for row in source_rows:
        warnings = []
        for field, value in row.items():
            if field.startswith("sl_") and isinstance(value, (int, float)) and value < 0:
                warnings.append({"field": field, "warning": "negative_source_value", "value": value})
        payload.append({
            "nguon_id": source_id,
            "dong_nguon": row["source_row"],
            "ma_quan_ly": row.get("ma_quan_ly"),
            "ten_quan_ly": row.get("ten_quan_ly"),
            "ma_hang": row.get("ma_hang"),
            "ten_hang": row.get("ten_vat_tu"),
            "ten_thuong_mai": row.get("ten_thuong_mai"),
            "so_quyet_dinh": row.get("so_quyet_dinh"),
            "dvt": row.get("dvt"),
            **{
                k: v for k, v in row.items()
                if k.startswith("sl_") or k.startswith("thoi_gian_") or k == "nha_cung_cap"
            },
            "canh_bao_chat_luong": warnings,
        })
    api.insert_ignore(
        "kha_dung_hop_dong_ma_hang",
        "nguon_id,dong_nguon",
        payload,
    )
    return {"nguon_id": source_id, "rows": len(payload)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--apply-snapshot", action="store_true")
    parser.add_argument(
        "--as-of-date",
        help="Ngày chốt số nghiệp vụ YYYY-MM-DD; để trống nếu file không ghi rõ.",
    )
    args = parser.parse_args()
    source = args.xlsx.resolve()
    if not source.exists():
        sys.exit(f"Không tìm thấy: {source}")

    staging_url = os.environ.get("SUPABASE_STAGING_URL", "").rstrip("/")
    staging_key = os.environ.get("SUPABASE_STAGING_SERVICE_ROLE_KEY", "")
    production_url = os.environ.get("SUPABASE_URL", "").rstrip("/")
    if not staging_url or not staging_key:
        sys.exit("Thiếu SUPABASE_STAGING_URL / SUPABASE_STAGING_SERVICE_ROLE_KEY.")
    if staging_url == production_url:
        sys.exit("DỪNG: URL staging trùng production.")

    digest = sha256(source)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    report_dir = BACKUP_ROOT / f"import-thoi-gian-su-dung-{stamp}-{digest[:10]}"
    report_dir.mkdir(parents=True, exist_ok=False)

    source_rows, excluded = parse(source)
    api = Rest(staging_url, staging_key)
    groups = api.fetch_all("nhom_ky_thuat", "ma_quan_ly,ten_quan_ly,created_at,updated_at")
    items = api.fetch_all("vat_tu", "*")
    plan = analyze(source_rows, groups, items)

    write_json(report_dir / "source_rows.json", source_rows)
    write_json(report_dir / "excluded_rows.json", excluded)
    write_json(report_dir / "backup_nhom_ky_thuat.json", groups)
    write_json(report_dir / "backup_vat_tu.json", items)
    write_json(
        report_dir / "plan.json",
        {
            "source_file": source.name,
            "source_sha256": digest,
            "staging_url": staging_url,
            "apply_requested": args.apply,
            "apply_snapshot_requested": args.apply_snapshot,
            "as_of_date": args.as_of_date,
            "excluded": excluded,
            "counts": plan["counts"],
            "negative_fields": plan["negative_fields"],
            "ambiguous_groups": plan["ambiguous_groups"],
            "source_ambiguities": plan["source_ambiguities"],
            "conflicts": plan["conflicts"],
            "new_groups": plan["new_groups"],
            "new_items": plan["new_items"],
            "patches": plan["patches"],
        },
    )

    print(json.dumps(plan["counts"], ensure_ascii=False, indent=2))
    print(f"Excluded source rows: {len(excluded)}")
    print(f"Report/backup: {report_dir}")
    if not args.apply and not args.apply_snapshot:
        print("DRY-RUN: chưa ghi staging. Chạy lại với --apply và/hoặc --apply-snapshot.")
        return

    if args.apply:
        api.insert_ignore("nhom_ky_thuat", "ma_quan_ly", plan["new_groups"])
        api.insert_ignore("vat_tu", "ma_hang", plan["new_items"])

    errors = []
    if args.apply:
        with ThreadPoolExecutor(max_workers=8) as pool:
            jobs = {
                pool.submit(api.patch_one, "vat_tu", "ma_hang", patch["ma_hang"], patch["changes"]): patch
                for patch in plan["patches"]
            }
            for future in as_completed(jobs):
                try:
                    future.result()
                except Exception as exc:
                    errors.append({"patch": jobs[future], "error": str(exc)})
    if errors:
        write_json(report_dir / "apply_errors.json", errors)
        raise RuntimeError(f"Có {len(errors)} PATCH lỗi; xem apply_errors.json.")

    groups_after = api.fetch_all("nhom_ky_thuat", "ma_quan_ly,ten_quan_ly")
    items_after = api.fetch_all("vat_tu", "*")
    after_groups = {row["ma_quan_ly"]: row for row in groups_after}
    after_items = {row["ma_hang"]: row for row in items_after}
    verify_errors = []
    if args.apply:
        for row in plan["new_groups"]:
            if row["ma_quan_ly"] not in after_groups:
                verify_errors.append({"missing_group": row})
        for row in plan["new_items"]:
            if row["ma_hang"] not in after_items:
                verify_errors.append({"missing_item": row})
        for patch in plan["patches"]:
            current = after_items.get(patch["ma_hang"], {})
            for key, expected in patch["changes"].items():
                if key == "updated_at":
                    continue
                if current.get(key) != expected:
                    verify_errors.append(
                        {
                            "ma_hang": patch["ma_hang"],
                            "field": key,
                            "expected": expected,
                            "actual": current.get(key),
                        }
                    )
    result = {
        "applied_at": datetime.now(timezone.utc).isoformat(),
        "source_sha256": digest,
        "groups_before": len(groups),
        "groups_after": len(groups_after),
        "items_before": len(items),
        "items_after": len(items_after),
        "planned": plan["counts"],
        "verify_errors": verify_errors,
    }
    if args.apply_snapshot:
        result["snapshot"] = apply_snapshot(
            api, source, digest, source_rows, args.as_of_date
        )
    write_json(report_dir / "result.json", result)
    if verify_errors:
        raise RuntimeError(f"Đối soát có {len(verify_errors)} lỗi; xem result.json.")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    print("APPLY + VERIFY: PASS")


if __name__ == "__main__":
    main()
