#!/usr/bin/env python3
"""Dựng bộ dữ liệu demo RHM từ các workbook nguồn.

Nguồn:
  - so luong su dung full.xlsx
  - RHM_tonghop_chaogia_3.7.xlsx
  - backend/thong tin vat tu y te tieu hao.xlsx

Đầu ra:
  - RHM/data/seed.json

Script này chỉ cần chạy lại khi workbook nguồn thay đổi. Web localhost không
đọc trực tiếp các file Excel lớn, nhờ vậy khởi động nhanh và không cần cài
openpyxl ở máy chạy demo.
"""

from __future__ import annotations

import json
import math
import re
import statistics
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


ROOT = Path(__file__).resolve().parents[2]
APP_ROOT = Path(__file__).resolve().parents[1]
USAGE_FILE = ROOT / "so luong su dung full.xlsx"
TENDER_FILE = ROOT / "RHM_tonghop_chaogia_3.7.xlsx"
CATALOG_FILE = ROOT / "backend" / "thong tin vat tu y te tieu hao.xlsx"
OUTPUT_FILE = APP_ROOT / "data" / "seed.json"

RHM_UNIT = "Khoa Phẫu thuật hàm mặt răng hàm mặt"
HISTORY_START = date(2024, 1, 1)
HISTORY_END = date(2026, 6, 30)
D12_START = date(2025, 7, 1)
DEFAULT_HORIZON = 18
DEFAULT_K_AB = 1.20
DEFAULT_K_C = 2.90
ABC_CUTOFF = 0.95


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_number(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def clean_number(value: float | None) -> int | float | None:
    if value is None:
        return None
    rounded = round(value)
    return int(rounded) if abs(value - rounded) < 1e-9 else round(value, 4)


def cell(row: tuple[Any, ...], column: str) -> Any:
    index = column_index_from_string(column) - 1
    return row[index] if index < len(row) else None


def split_his_codes(values: list[Any]) -> list[str]:
    codes: list[str] = []
    for value in values:
        text = as_text(value)
        if not text:
            continue
        for part in re.split(r"[\n,;/]+", text):
            part = part.strip()
            if re.fullmatch(r"\d+(?:\.0+)?", part):
                code = str(int(float(part)))
                if code not in codes:
                    codes.append(code)
    return codes


def month_range(start: date, end: date) -> list[str]:
    result: list[str] = []
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        result.append(f"{year:04d}-{month:02d}")
        month += 1
        if month == 13:
            year += 1
            month = 1
    return result


MONTHS = month_range(HISTORY_START, HISTORY_END)


def normalize_unit(value: str) -> str:
    text = unicodedata.normalize("NFD", value.lower())
    text = "".join(character for character in text if unicodedata.category(character) != "Mn")
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    aliases = {
        "gam": "gram",
        "g": "gram",
        "gram": "gram",
        "kg": "kg",
        "ky": "kg",
        "ml": "ml",
        "lit": "litre",
        "l": "litre",
        "cai": "count",
        "cay": "count",
        "que": "count",
        "mieng": "count",
        "chiec": "count",
        "soi": "count",
        "vien": "count",
        "thanh": "count",
        "te": "count",
        "lan": "count",
        "hop": "box",
        "goi": "pack",
        "chai": "bottle",
        "lo": "bottle",
        "tui": "bag",
        "bo": "set",
        "ong": "tube",
        "tuyp": "tube",
        "cuon": "roll",
        "vi": "blister",
    }
    return aliases.get(text, text)


def normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFD", value.lower())
    return "".join(character for character in text if unicodedata.category(character) != "Mn")


def conversion_factor(
    raw_unit: str,
    tender_unit: str,
    packaging: str,
) -> float | None:
    """Số đơn vị HIS tương ứng một ĐVT thầu."""

    raw = normalize_unit(raw_unit)
    tender = normalize_unit(tender_unit)
    if not raw or not tender:
        return None
    if raw == tender:
        return 1.0
    if raw == "count" and tender == "count":
        return 1.0
    if raw in {"gram", "kg"} and tender in {"gram", "kg"}:
        return 1000.0 if raw == "gram" and tender == "kg" else 0.001
    if raw in {"ml", "litre"} and tender in {"ml", "litre"}:
        return 1000.0 if raw == "ml" and tender == "litre" else 0.001

    normalized_packaging = normalize_text(packaging).replace(",", ".")
    # Quy cách nhiều lớp: Hộp/10 vỉ x 5 cái.
    if raw == "count":
        layered = re.search(
            r"(\d+(?:\.\d+)?)\s*vi\s*(?:x|×)\s*(\d+(?:\.\d+)?)\s*"
            r"(?:cai|cay|que|mieng|chiec|soi|vien|thanh)",
            normalized_packaging,
        )
        if layered:
            return float(layered.group(1)) * float(layered.group(2))
        count_match = re.search(
            r"(\d+(?:\.\d+)?)\s*(?:cai|cay|que|mieng|chiec|soi|vien|thanh)",
            normalized_packaging,
        )
        if count_match:
            return float(count_match.group(1))

    mass_match = re.search(r"(\d+(?:\.\d+)?)\s*(kg|g|gram|gam)\b", normalized_packaging)
    if mass_match and raw in {"gram", "kg"}:
        amount = float(mass_match.group(1))
        unit = mass_match.group(2)
        grams = amount * 1000 if unit == "kg" else amount
        return grams if raw == "gram" else grams / 1000

    volume_match = re.search(r"(\d+(?:\.\d+)?)\s*(ml|l|lit)\b", normalized_packaging)
    if volume_match and raw in {"ml", "litre"}:
        amount = float(volume_match.group(1))
        unit = volume_match.group(2)
        millilitres = amount * 1000 if unit in {"l", "lit"} else amount
        return millilitres if raw == "ml" else millilitres / 1000
    return None


def read_catalog() -> dict[str, dict[str, Any]]:
    workbook = load_workbook(CATALOG_FILE, read_only=True, data_only=True)
    worksheet = workbook["Sheet1"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [as_text(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = as_text(row[positions["Mã hàng"]])
        if not code:
            continue
        result[code] = {
            "package": as_text(row[positions["Gói"]]),
            "name": as_text(row[positions["Tên vật tư"]]),
            "unit": as_text(row[positions["ĐVT"]]),
            "specification": as_text(row[positions["Tiêu chí kỹ thuật"]]),
            "reference_product": as_text(row[positions["Tên thương mại"]]),
            "reference_model": as_text(row[positions["Ký mã hiệu"]]),
            "manufacturer": as_text(row[positions["Hãng"]]),
            "country": as_text(row[positions["Nước sản xuất"]]),
        }
    return result


def read_usage() -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(USAGE_FILE, read_only=True, data_only=True)
    worksheet = workbook["Export"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [as_text(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}

    by_code: dict[str, dict[str, Any]] = {}
    source_rows = 0
    missing_management_rows = 0

    for row in rows:
        unit_name = as_text(row[positions["Đơn vị"]])
        if unit_name != RHM_UNIT:
            continue
        source_rows += 1
        code = as_text(row[positions["Mã hàng"]])
        quantity = as_number(row[positions["Số lượng"]])
        raw_date = row[positions["Ngày"]]
        if not code or quantity is None or not isinstance(raw_date, (date, datetime)):
            continue
        used_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        if not (HISTORY_START <= used_date <= HISTORY_END):
            continue

        management_code = as_text(row[positions["Mã quản lý"]])
        if not management_code:
            missing_management_rows += 1

        record = by_code.setdefault(
            code,
            {
                "code": code,
                "names": Counter(),
                "management_codes": Counter(),
                "management_names": Counter(),
                "units": Counter(),
                "monthly": {month: 0.0 for month in MONTHS},
                "monthly_by_unit": defaultdict(
                    lambda: {month: 0.0 for month in MONTHS}
                ),
                "first_used": used_date,
                "last_used": used_date,
            },
        )
        record["names"][as_text(row[positions["Tên vật tư"]])] += 1
        record["management_codes"][management_code] += 1
        record["management_names"][as_text(row[positions["Tên quản lý"]])] += 1
        record["units"][as_text(row[positions["ĐVT"]])] += 1
        record["monthly"][used_date.strftime("%Y-%m")] += quantity
        record["monthly_by_unit"][as_text(row[positions["ĐVT"]])][
            used_date.strftime("%Y-%m")
        ] += quantity
        record["first_used"] = min(record["first_used"], used_date)
        record["last_used"] = max(record["last_used"], used_date)

    normalized: dict[str, dict[str, Any]] = {}
    for code, record in by_code.items():
        monthly = {key: clean_number(value) or 0 for key, value in record["monthly"].items()}
        management_code = record["management_codes"].most_common(1)[0][0]
        normalized[code] = {
            "code": code,
            "name": record["names"].most_common(1)[0][0],
            "management_code": management_code,
            "management_name": record["management_names"].most_common(1)[0][0],
            "unit": record["units"].most_common(1)[0][0],
            "unit_variants": [key for key, _ in record["units"].most_common() if key],
            "monthly": monthly,
            "monthly_by_unit": {
                unit: {
                    key: clean_number(value) or 0
                    for key, value in unit_monthly.items()
                }
                for unit, unit_monthly in record["monthly_by_unit"].items()
            },
            "first_used": record["first_used"].isoformat(),
            "last_used": record["last_used"].isoformat(),
        }

    quality = {
        "source_rows": source_rows,
        "distinct_his_codes": len(normalized),
        "missing_management_rows": missing_management_rows,
        "period_start": HISTORY_START.isoformat(),
        "period_end": HISTORY_END.isoformat(),
    }
    return normalized, quality


def read_old_price_map() -> dict[str, float]:
    """Lấy giá kế hoạch cũ làm fallback cho vài dòng chưa có báo giá mới."""

    workbook = load_workbook(TENDER_FILE, read_only=True, data_only=True)
    worksheet = workbook["raw"]
    result: dict[str, float] = {}
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number < 7:
            continue
        codes = split_his_codes([cell(row, key) for key in ("C", "D", "E", "F")])
        price = (
            as_number(cell(row, "AQ"))
            or as_number(cell(row, "AO"))
            or as_number(cell(row, "AN"))
        )
        if not price or price <= 0:
            continue
        for code in codes:
            result.setdefault(code, price)
    return result


def tender_monthly_fallback(row: tuple[Any, ...]) -> dict[str, int | float]:
    monthly = {month: 0 for month in MONTHS}
    month_columns_2025 = [
        "AF", "AG", "AH", "AI", "AJ", "AK",
        "AL", "AM", "AN", "AO", "AP", "AQ",
    ]
    month_columns_2026 = ["AR", "AS", "AT", "AU", "AV"]
    for month_number, column in enumerate(month_columns_2025, start=1):
        monthly[f"2025-{month_number:02d}"] = clean_number(as_number(cell(row, column))) or 0
    for month_number, column in enumerate(month_columns_2026, start=1):
        monthly[f"2026-{month_number:02d}"] = clean_number(as_number(cell(row, column))) or 0
    return monthly


def merge_monthly(
    codes: list[str],
    usage: dict[str, dict[str, Any]],
    fallback: dict[str, int | float],
) -> tuple[dict[str, int | float], str]:
    merged = {month: 0.0 for month in MONTHS}
    matched = False
    for code in codes:
        if code not in usage:
            continue
        matched = True
        for month in MONTHS:
            merged[month] += float(usage[code]["monthly"].get(month, 0) or 0)
    if not matched and any(float(value or 0) > 0 for value in fallback.values()):
        return fallback, "Workbook RHM (fallback do chưa khớp mã HIS)"
    return (
        {month: clean_number(value) or 0 for month, value in merged.items()},
        "Lịch sử xuất dùng HIS",
    )


def converted_tender_monthly(
    codes: list[str],
    usage: dict[str, dict[str, Any]],
    tender_monthly: dict[str, int | float],
    packaging: str,
    tender_unit: str,
) -> tuple[dict[str, int | float], str, str, list[str]]:
    """Ưu tiên số đã quy đổi trong workbook RHM và nối tháng 06/2026 từ HIS."""

    converted = {month: 0.0 for month in MONTHS}
    unresolved_series = {month: 0.0 for month in MONTHS}
    conversions: list[str] = []
    unresolved_units: list[str] = []
    source_units: set[str] = set()

    for code in codes:
        record = usage.get(code)
        if not record:
            continue
        for raw_unit, series in record.get("monthly_by_unit", {}).items():
            source_units.add(raw_unit)
            factor = conversion_factor(raw_unit, tender_unit, packaging)
            if factor is None or factor <= 0:
                unresolved_units.append(f"{code}: {raw_unit} → {tender_unit}")
                for month in MONTHS:
                    unresolved_series[month] += float(series.get(month, 0) or 0)
                continue
            if abs(factor - 1) > 1e-9:
                conversions.append(f"{code}: chia {factor:g} {raw_unit}/{tender_unit}")
            for month in MONTHS:
                converted[month] += float(series.get(month, 0) or 0) / factor

    controlled_months = [
        month
        for month in MONTHS
        if month.startswith("2025-") or "2026-01" <= month <= "2026-05"
    ]
    if any(unresolved_series.values()):
        raw_overlap = sum(unresolved_series[month] for month in controlled_months)
        tender_overlap = sum(float(tender_monthly[month] or 0) for month in controlled_months)
        known_overlap = sum(converted[month] for month in controlled_months)
        residual_tender = max(tender_overlap - known_overlap, 0)
        if raw_overlap > 0 and residual_tender > 0:
            inferred_factor = raw_overlap / residual_tender
            for month in MONTHS:
                converted[month] += unresolved_series[month] / inferred_factor
            conversions.append(
                f"Hệ số suy ra từ kỳ trùng: chia {inferred_factor:.4g}"
            )
            unresolved_units = []

    # Workbook RHM đã quy đổi theo ĐVT mời thầu cho 2025 và 01–05/2026.
    for month in controlled_months:
        converted[month] = float(tender_monthly.get(month, 0) or 0)

    monthly = {
        month: clean_number(value) or 0
        for month, value in converted.items()
    }
    source = "Workbook RHM đã quy đổi + HIS tháng 06/2026"
    note_parts = []
    if source_units:
        note_parts.append(f"ĐVT HIS: {', '.join(sorted(source_units))}")
    if conversions:
        note_parts.append("; ".join(dict.fromkeys(conversions)))
    note = " · ".join(note_parts)
    return monthly, source, note, sorted(set(unresolved_units))


def usage_metrics(monthly: dict[str, int | float]) -> dict[str, Any]:
    d12 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS if "2025-07" <= month <= "2026-06")
    use_2024 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS if month.startswith("2024-"))
    use_2025 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS if month.startswith("2025-"))
    use_2026 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS if month.startswith("2026-"))
    active_months = sum(1 for value in monthly.values() if float(value or 0) > 0)
    recent_6 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS[-6:])
    previous_6 = sum(float(monthly.get(month, 0) or 0) for month in MONTHS[-12:-6])
    trend = None if previous_6 == 0 else (recent_6 - previous_6) / previous_6
    return {
        "d12": clean_number(d12) or 0,
        "use_2024": clean_number(use_2024) or 0,
        "use_2025": clean_number(use_2025) or 0,
        "use_2026_h1": clean_number(use_2026) or 0,
        "active_months": active_months,
        "trend_6m": None if trend is None else round(trend, 4),
    }


def infer_reason_type(text: str) -> str:
    normalized = text.lower()
    if any(term in normalized for term in ("kỹ thuật mới", "danh mục mới", "vật tư mới")):
        return "ky_thuat_moi"
    if any(term in normalized for term in ("rớt thầu", "không trúng thầu", "thiếu hàng")):
        return "dut_hang"
    if any(term in normalized for term in ("cấp cứu", "cứu", "không có thay thế")):
        return "thiet_yeu"
    if any(term in normalized for term in ("giảm", "ngưng", "ít sử dụng")):
        return "giam_nhu_cau"
    if any(term in normalized for term in ("tăng", "số ca", "nhu cầu", "tần suất")):
        return "tang_nhu_cau"
    return "khac"


def read_tender(
    usage: dict[str, dict[str, Any]],
    catalog: dict[str, dict[str, Any]],
    old_prices: dict[str, float],
) -> tuple[list[dict[str, Any]], set[str]]:
    workbook = load_workbook(TENDER_FILE, read_only=True, data_only=True)
    worksheet = workbook["TRÌNH KÝ"]
    items: list[dict[str, Any]] = []
    mapped_codes: set[str] = set()

    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number < 8:
            continue
        order = as_number(cell(row, "A"))
        item_name = as_text(cell(row, "U")) or as_text(cell(row, "Q"))
        if order is None and not item_name:
            continue

        his_codes = split_his_codes([cell(row, key) for key in ("D", "E", "F", "G", "H")])
        mapped_codes.update(his_codes)
        catalog_rows = [catalog[code] for code in his_codes if code in catalog]
        primary_catalog = catalog_rows[0] if catalog_rows else {}
        fallback_monthly = tender_monthly_fallback(row)
        packaging = as_text(cell(row, "Y"))
        tender_unit = as_text(cell(row, "AA")) or primary_catalog.get("unit", "")
        monthly, usage_source, conversion_note, unresolved_units = converted_tender_monthly(
            his_codes,
            usage,
            fallback_monthly,
            packaging,
            tender_unit,
        )
        metrics = usage_metrics(monthly)
        sheet_use_2024 = as_number(cell(row, "AC"))
        sheet_use_2025 = as_number(cell(row, "AD"))
        sheet_use_2026 = as_number(cell(row, "AE"))
        if sheet_use_2024 is not None:
            metrics["use_2024"] = clean_number(sheet_use_2024) or 0
        if sheet_use_2025 is not None:
            metrics["use_2025"] = clean_number(sheet_use_2025) or 0
        if sheet_use_2026 is not None:
            metrics["use_2026_h1"] = clean_number(sheet_use_2026) or 0

        round_two_prices = [
            as_number(cell(row, key)) for key in ("CR", "CU", "CX", "DA", "DD")
        ]
        round_one_prices = [
            as_number(cell(row, key)) for key in ("BL", "BO", "BR", "BU", "BX")
        ]
        valid_round_two = [value for value in round_two_prices if value and value > 0]
        valid_round_one = [value for value in round_one_prices if value and value > 0]
        price_values = valid_round_two or valid_round_one
        price_source = (
            "Trung vị báo giá vòng 2"
            if valid_round_two
            else "Trung vị báo giá vòng 1"
            if valid_round_one
            else ""
        )
        if price_values:
            reference_price = statistics.median(price_values)
        else:
            old_values = [old_prices[code] for code in his_codes if code in old_prices]
            reference_price = statistics.median(old_values) if old_values else None
            if old_values:
                price_source = "Giá kế hoạch/trúng thầu 2024–2025 (fallback)"

        current_plan = as_number(cell(row, "BA"))
        if current_plan is None:
            current_plan = as_number(cell(row, "AX"))
        current_option = as_number(cell(row, "BB"))
        if current_option is None:
            current_option = as_number(cell(row, "AY"))
        current_reason = as_text(cell(row, "BC")) or as_text(cell(row, "AZ"))

        management_code = as_text(cell(row, "K"))
        technical_code = as_text(cell(row, "M"))
        data_flags: list[str] = []
        if not his_codes:
            data_flags.append("Chưa có mã HIS")
        if not management_code:
            data_flags.append("Chưa có mã quản lý")
        if reference_price is None:
            data_flags.append("Thiếu đơn giá")
        if not any(code in usage for code in his_codes):
            data_flags.append("Chưa khớp lịch sử HIS")
        if metrics["d12"] == 0:
            data_flags.append("Không có sử dụng 12 tháng")
        if unresolved_units:
            data_flags.append("Chưa đủ hệ số quy đổi ĐVT")

        item = {
            "id": f"T-{int(order) if order is not None else row_number}",
            "order": int(order) if order is not None else row_number,
            "scope": "chuyen_khoa",
            "scope_label": "Danh mục thầu chuyên khoa",
            "his_codes": his_codes,
            "management_code": management_code,
            "management_name": as_text(cell(row, "L")),
            "technical_code": technical_code,
            "technical_name": as_text(cell(row, "N")),
            "item_name": item_name,
            "unit": tender_unit,
            "package": primary_catalog.get("package", "Răng Hàm Mặt"),
            "packaging": packaging,
            "specification": as_text(cell(row, "W")) or as_text(cell(row, "V")),
            "reference_product": as_text(cell(row, "BE"))
            or primary_catalog.get("reference_product", ""),
            "reference_model": as_text(cell(row, "BF"))
            or primary_catalog.get("reference_model", ""),
            "manufacturer": as_text(cell(row, "BG"))
            or primary_catalog.get("manufacturer", ""),
            "country": as_text(cell(row, "BH")) or primary_catalog.get("country", ""),
            "reference_price": clean_number(reference_price),
            "price_source": price_source,
            "monthly": monthly,
            "usage_source": usage_source,
            "conversion_note": conversion_note,
            "unresolved_conversions": unresolved_units,
            **metrics,
            "current_plan_qty": clean_number(current_plan),
            "current_option_qty": clean_number(current_option),
            "current_reason": current_reason,
            "reason_type": infer_reason_type(current_reason),
            "data_flags": data_flags,
        }
        items.append(item)

    return items, mapped_codes


def read_supplemental_usage(
    usage: dict[str, dict[str, Any]],
    mapped_codes: set[str],
    catalog: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for index, code in enumerate(sorted(set(usage) - mapped_codes), start=1):
        source = usage[code]
        catalog_row = catalog.get(code, {})
        metrics = usage_metrics(source["monthly"])
        package = catalog_row.get("package", "")
        scope = "dung_chung" if package == "Dùng chung" else "ngoai_danh_muc"
        scope_label = (
            "Vật tư dùng chung"
            if scope == "dung_chung"
            else "Có sử dụng, ngoài danh mục chuyên khoa"
        )
        data_flags = ["Chưa có trong danh mục thầu RHM", "Thiếu đơn giá"]
        if not source["management_code"]:
            data_flags.append("Chưa có mã quản lý")
        if len(source["unit_variants"]) > 1:
            data_flags.append("Nhiều đơn vị tính")

        items.append(
            {
                "id": f"H-{code}",
                "order": 1000 + index,
                "scope": scope,
                "scope_label": scope_label,
                "his_codes": [code],
                "management_code": source["management_code"],
                "management_name": source["management_name"],
                "technical_code": "",
                "technical_name": "",
                "item_name": catalog_row.get("name") or source["name"],
                "unit": source["unit"] or catalog_row.get("unit"),
                "package": package or "Chưa gán gói",
                "packaging": "",
                "specification": catalog_row.get("specification", ""),
                "reference_product": catalog_row.get("reference_product", ""),
                "reference_model": catalog_row.get("reference_model", ""),
                "manufacturer": catalog_row.get("manufacturer", ""),
                "country": catalog_row.get("country", ""),
                "reference_price": None,
                "price_source": "",
                "monthly": source["monthly"],
                "usage_source": "Lịch sử xuất dùng HIS",
                "conversion_note": "Giữ nguyên ĐVT HIS",
                "unresolved_conversions": [],
                **metrics,
                "current_plan_qty": None,
                "current_option_qty": None,
                "current_reason": "",
                "reason_type": "theo_lich_su",
                "data_flags": data_flags,
            }
        )
    return items


def assign_abc(items: list[dict[str, Any]]) -> None:
    groups: dict[str, dict[str, Any]] = {}
    for item in items:
        key = item["management_code"] or item["technical_code"] or item["id"]
        group = groups.setdefault(
            key,
            {"items": [], "quantity": 0.0, "value": 0.0, "price_complete": True},
        )
        group["items"].append(item)
        quantity = float(item["d12"] or 0)
        group["quantity"] += quantity
        if quantity > 0 and item["reference_price"] is None:
            group["price_complete"] = False
        if item["reference_price"] is not None:
            group["value"] += quantity * float(item["reference_price"])

    value_groups = [
        (key, group)
        for key, group in groups.items()
        if group["quantity"] > 0 and group["price_complete"] and group["value"] > 0
    ]
    quantity_groups = [
        (key, group)
        for key, group in groups.items()
        if group["quantity"] > 0 and not (group["price_complete"] and group["value"] > 0)
    ]

    assignments: dict[str, tuple[str, str]] = {}
    for subset, metric, method in (
        (value_groups, "value", "Giá trị tiền"),
        (quantity_groups, "quantity", "Số lượng (tạm do thiếu giá)"),
    ):
        total = sum(float(group[metric]) for _, group in subset)
        cumulative = 0.0
        for key, group in sorted(subset, key=lambda pair: pair[1][metric], reverse=True):
            share_before = cumulative / total if total else 1.0
            assignments[key] = ("A+B" if share_before < ABC_CUTOFF else "C", method)
            cumulative += float(group[metric])

    for key, group in groups.items():
        if group["quantity"] <= 0:
            abc_group, abc_method = "MỚI", "Chưa có D12"
        else:
            abc_group, abc_method = assignments[key]
        k = DEFAULT_K_AB if abc_group == "A+B" else DEFAULT_K_C if abc_group == "C" else 0
        for item in group["items"]:
            item["abc_group"] = abc_group
            item["abc_method"] = abc_method
            item["k_default"] = k
            item["estimated_annual_value"] = clean_number(
                float(item["d12"] or 0) * float(item["reference_price"] or 0)
            )
            formula = float(item["d12"] or 0) * DEFAULT_HORIZON / 12 * k
            item["formula_default"] = math.ceil(formula)
            current_plan = item["current_plan_qty"]
            item["initial_working_qty"] = (
                current_plan if current_plan is not None else math.ceil(formula)
            )
            working = float(item["initial_working_qty"] or 0)
            if formula == 0:
                delta = None if working == 0 else 1.0
            else:
                delta = (working - formula) / formula
            item["delta_vs_formula"] = None if delta is None else round(delta, 4)
            item["reason_required"] = bool(
                (formula == 0 and working > 0)
                or (formula > 0 and abs(working - formula) / formula > 0.30)
                or (formula > 0 and working == 0)
            )


def build_seed() -> dict[str, Any]:
    catalog = read_catalog()
    usage, usage_quality = read_usage()
    old_prices = read_old_price_map()
    tender_items, mapped_codes = read_tender(usage, catalog, old_prices)
    supplemental_items = read_supplemental_usage(usage, mapped_codes, catalog)
    items = tender_items + supplemental_items
    assign_abc(items)

    price_known = sum(1 for item in items if item["reference_price"] is not None)
    mapped_seen_codes = set(usage) & mapped_codes
    activity_monthly = {
        month: sum(
            float(record["monthly"].get(month, 0) or 0) > 0
            for record in usage.values()
        )
        for month in MONTHS
    }

    return {
        "meta": {
            "title": "Bảng đề xuất số lượng đi thầu — Khoa RHM",
            "unit": RHM_UNIT,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_files": [
                USAGE_FILE.name,
                TENDER_FILE.name,
                str(CATALOG_FILE.relative_to(ROOT)),
                "phan-tich-cong-thuc/bao-cao-backtest.md",
                "cong-thuc-dat-so-luong-VTYT.md",
            ],
            "history_start": HISTORY_START.isoformat(),
            "history_end": HISTORY_END.isoformat(),
            "default_period_from": "2027-01",
            "default_period_to": "2028-06",
            "default_horizon_months": DEFAULT_HORIZON,
            "formula_version": "UMC VTYT v1.0 — 28/07/2026",
            "notes": [
                "Số lượng đề xuất hiện hữu lấy từ sheet TRÌNH KÝ của workbook RHM.",
                "Công thức là hàng rào mức phục vụ, không phải dự báo điểm.",
                "Chưa có tồn kho, hàng đang về và hạn dùng nên chưa thể chốt nhu cầu ròng.",
            ],
        },
        "quality": {
            **usage_quality,
            "tender_lines": len(tender_items),
            "supplemental_lines": len(supplemental_items),
            "mapped_his_codes": len(mapped_codes),
            "mapped_his_codes_seen": len(mapped_seen_codes),
            "mapped_his_code_share": round(
                len(mapped_seen_codes) / len(usage), 4
            )
            if usage
            else 0,
            "price_known_lines": price_known,
            "price_coverage_lines": round(price_known / len(items), 4) if items else 0,
        },
        "settings": {
            "period_from": "2027-01",
            "period_to": "2028-06",
            "horizon_months": DEFAULT_HORIZON,
            "transition_delay_months": 0,
            "r_adjust": 1.0,
            "k_ab": DEFAULT_K_AB,
            "k_c": DEFAULT_K_C,
            "abc_cutoff": ABC_CUTOFF,
        },
        "activity_monthly": activity_monthly,
        "items": items,
    }


def main() -> None:
    seed = build_seed()
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(
        json.dumps(seed, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Đã tạo {OUTPUT_FILE}")
    print(
        f"{len(seed['items'])} dòng: "
        f"{seed['quality']['tender_lines']} chuyên khoa + "
        f"{seed['quality']['supplemental_lines']} ngoài danh mục."
    )


if __name__ == "__main__":
    main()
