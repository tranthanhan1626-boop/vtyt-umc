"""Kiểm nhanh: hình nào lấn dải chân (y > 5.40 in), chữ nào < 12pt, ảnh có méo không."""
import sys
from pptx import Presentation
from pptx.util import Emu
E = 914400
f = sys.argv[1] if len(sys.argv) > 1 else "../../../Test và hướng dẫn 10-10-2026/HuongDan_VTYT_10-2026.pptx"
p = Presentation(f)
bad = 0
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        b = (sh.top + sh.height) / E
        r = (sh.left + sh.width) / E
        if (b > 5.40 or r > 9.62 or sh.left < 0 or sh.top < 0) and i > 1:
            print(f"trang {i}: '{sh.name}' lấn mép (đáy {b:.2f}, phải {r:.2f}) {sh.text_frame.text[:40] if sh.has_text_frame else ''}"); bad += 1
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None and run.font.size.pt < 12:
                        print(f"trang {i}: chữ {run.font.size.pt}pt: {run.text[:30]}"); bad += 1
        if sh.shape_type == 13:
            from PIL import Image
            import io
            im = Image.open(io.BytesIO(sh.image.blob))
            W, H = im.size
            cw = W * (1 - sh.crop_left - sh.crop_right)
            ch = H * (1 - sh.crop_top - sh.crop_bottom)
            ratio_img = cw / ch
            ratio_box = sh.width / sh.height
            if abs(ratio_img / ratio_box - 1) > 0.01:
                print(f"trang {i}: ảnh méo {ratio_img:.3f} vs {ratio_box:.3f}"); bad += 1
            if sh.crop_left + sh.crop_right > 0.185 or sh.crop_top + sh.crop_bottom > 0.185:
                print(f"trang {i}: ảnh cắt >18% ({W}x{H})")
print("Số lỗi:", bad)
