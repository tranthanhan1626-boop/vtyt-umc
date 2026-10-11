"""Dựng 2 ảnh phóng to bản Excel CHÍNH THỨC cho slide (11/10/2026).
Ẩn 8 cột trống + tạm ẩn 5 dòng, chia 2 nửa cột, in vừa 1 tờ, cắt sát nội dung, xuất PNG.
Chạy: ../slide/.venv/bin/python dung_excel_phong_to.py  (trong thư mục ket-qua)"""
import json, math, subprocess, openpyxl, pymupdf
from openpyxl.utils import get_column_letter as L

SRC = "excel_chinh_thuc_dung_chung.xlsx"
STT_GIU = {1, 2, 7, 9, 10}
PHAN = {
    "phan1": ({"B": 30, "C": 6, "I": 15, "J": 18, "L": 62, "N": 9,
               "O": 11, "P": 11, "Q": 11, "R": 11, "S": 11, "T": 11, "U": 11, "V": 11}, False),
    "phan2": ({"B": 30, "C": 6, "W": 12, "X": 14, "Y": 12, "Z": 34, "AA": 32, "AB": 11,
               "AC": 24, "AD": 12, "AE": 12, "AF": 12, "AG": 12}, True),
}


def so_dong(txt, rong):
    if txt in (None, ""):
        return 1
    return sum(max(1, math.ceil(len(d) / max(1, rong * 1.05))) for d in str(txt).split("\n"))


for ten, (cot, gon) in PHAN.items():
    wb = openpyxl.load_workbook(SRC)
    ws = wb.active
    for c in range(1, ws.max_column + 1):
        l = L(c)
        if l in cot:
            ws.column_dimensions[l].width = cot[l]
        else:
            ws.column_dimensions[l].hidden = True
    ws.row_dimensions[5].height = 66  # tiêu đề cột 3–4 dòng chữ, khỏi bị cắt
    for r in range(6, ws.max_row + 1):
        if ws.cell(r, 3).value not in STT_GIU:
            ws.row_dimensions[r].hidden = True
        elif gon:  # nửa 2 không có cột mô tả → đặt chiều cao theo cột đang hiện
            n = max(so_dong(ws.cell(r, c).value, cot[L(c)]) for c in range(1, ws.max_column + 1) if L(c) in cot)
            ws.row_dimensions[r].height = 13.5 * n + 4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    wb.save(f"{ten}.xlsx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", f"{ten}.xlsx"], capture_output=True)
    p = pymupdf.open(f"{ten}.pdf")[0]
    bb = None
    for b in p.get_text("blocks"):
        r = pymupdf.Rect(b[:4]); bb = r if bb is None else bb | r
    for dr in p.get_drawings():
        bb |= dr["rect"]
    bb = pymupdf.Rect(bb.x0 - 4, bb.y0 - 4, bb.x1 + 4, bb.y1 + 4)
    pix = p.get_pixmap(dpi=300, clip=bb)
    pix.save(f"excel_{ten}.png")
    k = 300 / 72  # pt → px của ảnh
    vi_tri = {}
    for tu in ("Số lượng", "TỔNG TOÀN", "Khoa Ngoại", "Bộ thông tiểu dùng 1 lần", "Mô tả và đặc tính",
               "Đơn vị", "Giải trình đề xuất", "BẢN CHÍNH THỨC"):
        rs = p.search_for(tu, clip=bb)
        if rs:
            r = rs[0]
            vi_tri[tu] = [round((r.x0 - bb.x0) * k), round((r.y0 - bb.y0) * k),
                          round((r.x1 - bb.x0) * k), round((r.y1 - bb.y0) * k)]
    vi_tri["_kich_thuoc"] = [pix.width, pix.height]
    json.dump(vi_tri, open(f"excel_{ten}_vi_tri.json", "w"), ensure_ascii=False)
    print(ten, pix.width, pix.height, round(pix.width / pix.height, 2))
